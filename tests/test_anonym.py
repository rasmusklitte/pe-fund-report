"""Den anonymiserede udgave: navne, identitet og beløb (anonym.py og slides.beloebsskala)."""
from datetime import date

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from pe_fund_report import anonym, cli, slides
from pe_fund_report.data import Note
from pe_fund_report.deck import L_TEKST, _ny_slide, _ryd_slides
from pe_fund_report.paths import MASTER


def _raekke(navn, type_="CURRENT"):
    return dict(navn=navn, type=type_, sektor="Niche hobby consumables", dagsvaerdi=1000.0)


def _fond(kode, navn, segment, raekker, via=None):
    beh = dict(dato=date(2026, 6, 30), kode=via and "UND1" or kode, raekker=raekker,
               type="selskaber", fakta={"NUMBER_OF_FUNDS": 26.0, "CREDIT_LINE_DRAWN": 5e7})
    if via:
        beh.update(via=via, via_kode="UND1")
    return dict(
        fund_code=kode, fund_name=navn, juridisk_navn=f"{navn} K/S", segment=segment,
        manager="Acme Capital Management A/S", logo=dict(billede=b"png"),
        beskrivelse=f"{navn} investerer i Army Painter.", virksomheder={"Army Painter": {}},
        benchmark="MSCI", nav_afstemning=dict(dele=[]),
        n=dict(nav=100.0, fondsstoerrelse=5e8, fonds_nav=4e8, andel=0.02),
        beholdninger=beh, beholdninger_primo={"kvartal": dict(beh, raekker=[_raekke("army painter")])},
        kilder=dict(nav=None),
        noter=[Note(f"{navn} har givet 1.234 tEUR i rabat og 17.574 DKK i rente (04.09.2023).",
                    "indbetalt")])


def _rapport():
    a = _fond("CATA3", "CataCap III", "Private equity",
              [_raekke("The Army Painter"), _raekke("Enwerdi")])
    b = _fond("PPPC2", "Petersen & Partners Private Credit II", "Private credit",
              [_raekke("ARMY PAINTER"), _raekke("Advent GPE X", "FUND")],
              via="Park Square Capital Partners V, SCSp")
    a["beholdninger"]["raekker"][0]["navn"] = "Army Painter"
    return dict(investor="MOTORTRAMP", navn="A/S Motortramp", fonde=[a, b],
                fordelinger={"segment": dict(grupper=[dict(navn="Private equity", fonde=[a])])},
                stoerste_positioner=dict(linjer=[dict(navn="Army Painter", fonde=["CataCap III"])]),
                dokumenter="df", kursopslag={})


def _alle_strenge(x):
    if isinstance(x, str):
        yield x
    elif isinstance(x, dict):
        for k, v in x.items():
            yield from _alle_strenge(k)
            yield from _alle_strenge(v)
    elif isinstance(x, list | tuple):
        for v in x:
            yield from _alle_strenge(v)


def test_anonymiser_fjerner_alle_navne():
    r = _rapport()
    ny, originaler = anonym.anonymiser(r, 2.0)
    tekst = "\n".join(_alle_strenge(ny)).lower()
    for navn in ("catacap", "petersen", "park square", "army painter", "enwerdi", "advent",
                 "motortramp", "acme", "cata3", "pppc2", "und1"):
        assert navn not in tekst, navn
    assert {"catacap iii", "army painter", "park square capital partners v"} <= set(originaler)
    # r selv er urørt; laek bruger de rigtige navne
    assert r["fonde"][0]["fund_name"] == "CataCap III"


def test_anonymiser_pseudonymer():
    ny, _ = anonym.anonymiser(_rapport(), 2.0)
    a, b = ny["fonde"]
    assert (ny["investor"], ny["navn"]) == ("Investor", "Investor")
    assert (a["fund_name"], b["fund_name"]) == ("Private Equity-fond 1", "Private Credit-fond 1")
    assert b["beholdninger"]["via"] == "Underliggende fond 1"
    # samme selskab i to fonde og i primo, uanset store og små bogstaver
    navne = [x["navn"] for x in a["beholdninger"]["raekker"] + b["beholdninger"]["raekker"]]
    assert navne == ["Selskab 1", "Selskab 2", "Selskab 1", "Underliggende fond 2"]
    assert a["beholdninger_primo"]["kvartal"]["raekker"][0]["navn"] == "Selskab 1"
    assert ny["stoerste_positioner"]["linjer"][0] == dict(navn="Selskab 1",
                                                          fonde=["Private Equity-fond 1"])


def test_anonymiser_fjerner_identitet_overalt():
    ny, _ = anonym.anonymiser(_rapport(), 2.0)
    for f in ny["fonde"]:
        assert f["logo"] is f["benchmark"] is f["nav_afstemning"] is None
        assert f["manager"] == "Forvalter A"     # begge fonde har samme forvalter
        assert f["virksomheder"] == {} and f["beskrivelse"] == anonym.BESKRIVELSE
        assert f["n"]["fondsstoerrelse"] is f["n"]["fonds_nav"] is f["n"]["andel"] is None
        assert f["n"]["nav"] == 100.0           # beløbene skaleres i visningen, ikke her
        assert "kilder" not in f
        assert f["beholdninger"]["fakta"] == {"CREDIT_LINE_DRAWN": 5e7}
        assert f["beholdninger"]["raekker"][0]["sektor"] is None
    # fonden i fordelingerne er samme objekt som i fonde
    assert ny["fordelinger"]["segment"]["grupper"][0]["fonde"][0] is ny["fonde"][0]
    assert "dokumenter" not in ny and "kursopslag" not in ny


def test_anonymiser_skalerer_beloeb_i_noter():
    ny, _ = anonym.anonymiser(_rapport(), 2.0)
    note = ny["fonde"][0]["noter"][0]
    assert isinstance(note, Note) and note.felt == "indbetalt"
    assert note == ("Private Equity-fond 1 har givet 2.468 tEUR i rabat og 35.148 DKK i rente "
                    "(04.09.2023).")


def test_beloebsskala():
    with slides.beloebsskala(2.0):
        assert slides.tusind(1_234_567) == "2.469"
        assert slides.million(1_500_000) == "3"
        assert slides.kroner(1_000) == "2.000"
        assert slides._kt(1_000) == 2
    assert slides.tusind(1_234_567) == "1.235" and slides._kt(1_000) == 1


def test_laek_finder_navne_i_tabel_diagram_og_noter():
    prs = Presentation(str(MASTER))
    _ryd_slides(prs)
    s = _ny_slide(prs, L_TEKST)
    tabel = s.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    tabel.cell(0, 1).text = "Selskab 1"
    assert anonym.laek(prs, ["army painter", "catacap iii"]) == []
    tabel.cell(0, 0).text = "The Army Painter"
    cd = CategoryChartData()
    cd.categories = ["CataCap III"]
    cd.add_series("x", [1])
    s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(2), cd)
    s.notes_slide.notes_text_frame.text = "Fra Motortramp"
    assert anonym.laek(prs, ["army painter", "catacap iii", "motortramp", "enwerdi"]) == [
        "army painter", "catacap iii", "motortramp"]


def test_cli_anonym(monkeypatch):
    kald = {}
    monkeypatch.setattr(cli, "byg", lambda *a, **k: kald.update(a=a, k=k))
    cli.main(["MOTORTRAMP", "--as-of", "2026-06-30", "--anonym"])
    assert kald["k"] == dict(med_data=True, anonym=True, med_supplerende=True)
    cli.main(["MOTORTRAMP", "--uden-supplerende"])
    assert kald["k"]["anonym"] is False and kald["k"]["med_supplerende"] is False


def test_kort_fondsnavn_beholder_loebenummer():
    assert slides._kort_fondsnavn("Private Credit-fond 1") == "Private Credit-fond 1"
    assert slides._kort_fondsnavn("Fund of funds 2") == "Fund of funds 2"
    assert slides._kort_fondsnavn("K Alternativ Private Equity 2020 K/S") == "KAPE 2020"
