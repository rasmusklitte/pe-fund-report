"""Afstemning af Motortramp-rapporten mod Aletas tal (data/aleta/MOTORTRAMP/; indlæsningen og
parringen af transaktionerne er pe_fund_report.aleta, som decket også bruger).

Aleta er et andet system med de samme fonde; udtrækket har én række pr. fond. afstem() sammenligner
rapportens nøgletal med Aletas, kontrol for kontrol. Aletas beløb er i DKK (undtagen nav, der er i
fondens valuta). Aletas udtræk er altid dateret den første i måneden efter rapportdatoen
(aleta_dato(): 30.06.2026 -> 01.07.2026), og Aleta omregner beholdningerne (NAV, tilsagn,
resttilsagn) med kursen den dag. Rapportens beløb i fondens valuta omregnes derfor med
Nationalbankens kurs på Aletas dato, før de sammenlignes; indbetalt og udloddet sammenlignes med
rapportens egen omregning (n["..._dkk"] fra data.omregn_til_dkk) med kursen på hver kapitalstrøms
egen dato, som Aleta gør. afstem_transaktioner() parrer Aletas transaktioner (*transactions_*.xlsx)
med rapportens kapitalstrømme, og forklar() bruger dem til årsagerne. test_aleta.py kører
kontrollerne og skriver afstemningsdokumentet (skriv()) med de samme linjer, så de altid er enige:

    python -m pytest tests/test_aleta.py     # testene + exports/aleta_afstemning_<dato>.xlsx
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pe_fund_report import aleta as pa
from pe_fund_report.aleta import (
    FONDE,
    KONTROLLER,
    TRANSAKTIONSSTATUS,
    _aletas_irr,
    _ens,
    _gp_konto,
    _kr,
    _uden_nav,
    aleta_dato,
)

INVESTOR = "MOTORTRAMP"
AS_OF = date(2026, 6, 30)
ALETA_DATO = aleta_dato(AS_OF)
# Aletas nøgletal og transaktioner (det nyeste udtræk) i data/aleta/MOTORTRAMP/
ALETA, TRANSAKTIONER = pa.find(INVESTOR, AS_OF) or (None, None)
# kontrollerne (KONTROLLER, _vurder, afstem) bor i pe_fund_report.aleta, så decket og testene
# giver samme status; her lægges forklaringerne og Excel-afstemningen ovenpå


def laes_aleta(sti=None):
    """Aletas rækker (pe_fund_report.aleta.laes_aleta) fra Motortramps udtræk."""
    return pa.laes_aleta(sti or ALETA)


def afstem(aleta, r, kurser_aleta=None, transaktioner=None):
    """pe_fund_report.aleta.afstem() pr. Aletas dato, med årsagen til hver forklarlig afvigelse i
    bemaerkning (forklar()); transaktioner er linjerne fra afstem_transaktioner(), som forklar()
    bruger til årsagerne, når de er givet."""
    ud = pa.afstem(aleta, r, ALETA_DATO, kurser_aleta)
    fonde = {f["fund_code"]: f for f in r["fonde"]}
    for a in aleta:
        kode = FONDE.get(a["securityCode"], a["securityCode"])
        if kode in fonde:
            forklar([x for x in ud if x["fond"] == kode], a, fonde[kode], kurser_aleta,
                    [t for t in transaktioner or [] if t["fond"] == kode])
    return ud


# --- transaktionerne ------------------------------------------------------------------------------

def laes_transaktioner(sti=None):
    """Aletas transaktioner (pe_fund_report.aleta.laes_transaktioner) fra det nyeste udtræk."""
    return pa.laes_transaktioner(sti or TRANSAKTIONER)


def afstem_transaktioner(transaktioner, r, as_of=AS_OF):
    """pe_fund_report.aleta.afstem_transaktioner pr. Motortramp-rapportens dato."""
    return pa.afstem_transaktioner(transaktioner, r, as_of)


# --- årsager til afvigelserne --------------------------------------------------------------------

def _pct(x):
    return f"{x * 100:.2f} %".replace(".", ",")


def _kurs(x):
    return f"{x:.4f}".replace(".", ",")


def _irr_dkk(f, nav_dkk, datoer=None):
    return pa._irr_dkk(f, nav_dkk, ALETA_DATO, datoer)


NAVNE = {"NAV DKK": "NAV", "Indbetalt DKK": "indbetalt", "Udloddet DKK": "udloddet"}


def _tx_tekst(t):
    """Én afvigende kapitalstrøm som tekst til en bemærkning."""
    ref = f" ({t['ref']})" if t["ref"] else ""
    if t["art"] == "kun i Aleta":
        return f"kun i Aleta: {_kr(t['aleta'])} pr. {t['aleta_dato']:%d.%m.%Y}"
    if t["art"] == "kun i rapporten":
        return f"kun i rapporten: {_kr(t['rapport'])} pr. {t['rapport_dato']:%d.%m.%Y}{ref}"
    return (f"pr. {t['rapport_dato']:%d.%m.%Y} har Aleta {_kr(t['aleta'])}, rapporten "
            f"{_kr(t['rapport'])}{ref}")


def _stroemme(antal):
    return f"{antal} strøm" if antal == 1 else f"{antal} strømme"


TYPER = {"Indbetalt DKK": "kald", "Udloddet DKK": "udlodning"}


def forklar(linjer, a, f, kurser_aleta=None, transaktioner=None):
    """Skriv årsagen til hver afvigelse i bemaerkning, hvor tallene viser den. linjer er én fonds
    linjer fra afstem(), a Aletas række, f rapportens fond og transaktioner fondens linjer fra
    afstem_transaktioner(). Kun årsager, der kan regnes efter; ellers står bemaerkning, som den
    var."""
    from pe_fund_report.data import xirr
    x = {lin["kontrol"]: lin for lin in linjer}
    afviger = {kontrol for kontrol, lin in x.items() if lin["status"] != "ok"}
    n, valuta = f["n"], f["valuta"]
    dkk = valuta == "DKK"
    k_aleta = 1.0 if dkk else ((kurser_aleta or {}).get(valuta) or (None,))[0]
    ca890, ca900, ca110 = (_gp_konto(f, kode) for kode in ("CA890", "CA900", "CA110"))
    foer_carry = ca900 is not None and _ens(a["nav"], ca890)
    tx = [t for t in transaktioner or [] if t["art"] != "efter rapportdatoen"]
    tx_afviger = [t for t in tx if t["status"] == "afviger"]
    tx_datoer = [t for t in tx if t["art"] == "anden dato"]

    def skriv(kontrol, tekst):
        lin = x[kontrol]
        lin["bemaerkning"] = "; ".join(t for t in (lin["bemaerkning"], tekst) if t)

    # NAV: Aletas er før carry, eller rapportens er forvalterens kapitalkonto
    for kontrol in ("NAV", "NAV DKK"):
        if kontrol not in afviger:
            continue
        if foer_carry:
            skriv(kontrol, f"Aleta bruger NAV før carried interest (kapitalkontoens CA890 "
                           f"{_kr(ca890)} {valuta}); rapporten NAV efter carry (CA900 "
                           f"{_kr(ca900)}), carry {_kr(ca890 - ca900)}")
        elif _ens(n["nav"], ca900):
            mere = a["nav"] - n["nav"]
            skriv(kontrol, f"Rapportens NAV er forvalterens kapitalkonto (CA900 {_kr(ca900)} "
                           f"{valuta}); Aleta har {_kr(abs(mere))} "
                           f"{'mere' if mere > 0 else 'mindre'}")

    # indbetalt og udloddet: udligningsrente, kurs eller forvalterens kapitalkonto
    udligning = (n.get("udligning") or 0) * (k_aleta or 0)
    for kontrol, noegle in (("Indbetalt DKK", "indbetalt"), ("Udloddet DKK", "udloddet")):
        lin = x[kontrol]
        if kontrol not in afviger or lin["afvigelse"] is None:
            continue
        forkert = [t for t in tx_afviger if t["type"] == TYPER[kontrol]]
        if forkert:
            skriv(kontrol, "Transaktionerne: " + "; ".join(_tx_tekst(t) for t in forkert[:3])
                  + (f" og {len(forkert) - 3} til" if len(forkert) > 3 else ""))
        omregnet = [t for t in tx if t["type"] == TYPER[kontrol] and t["afvigelse_dkk"] is not None
                    and t["status"] != "afviger"]
        if not dkk and omregnet and not forkert:
            t = max(omregnet, key=lambda t: abs(t["afvigelse_dkk"]))
            skriv(kontrol, f"Aleta omregner med sine egne kurser på sine egne datoer, fx "
                           f"{_kurs(t['aleta_kurs'])} pr. {t['aleta_dato']:%d.%m.%Y} mod "
                           f"Nationalbankens {_kurs(t['rapport_kurs'])} pr. "
                           f"{t['rapport_dato']:%d.%m.%Y}; i alt "
                           f"{_kr(sum(t['afvigelse_dkk'] for t in omregnet))} DKK over "
                           f"{_stroemme(len(omregnet))} (se fanen Transaktioner)")
            continue
        if noegle == "indbetalt" and udligning and _ens(lin["afvigelse"], udligning, 2):
            skriv(kontrol, f"Aleta regner udligningsrenten ({_kr(udligning)} DKK) ud af "
                           "indbetalt; rapporten regner den med")
        elif not dkk and n.get(noegle) and lin["kurs"] and not tx:
            skriv(kontrol, f"Svarer til en kurs på {_kurs(lin['aleta'] / n[noegle])} hos Aleta "
                           f"mod Nationalbankens {_kurs(lin['kurs'])} (vægtet over "
                           f"transaktionsdatoerne), hvis beløbet i {valuta} er det samme")
        elif noegle == "indbetalt" and _ens(n["indbetalt"], ca110):
            mere = -lin["afvigelse"]
            skriv(kontrol, f"Rapporten følger forvalterens kapitalkonto (CA110 {_kr(ca110)}); "
                           f"Aleta har {_kr(abs(mere))} {'mere' if mere > 0 else 'mindre'}")

    # resttilsagn er tilsagn - indbetalt hos Aleta
    rest, indb = x["Resttilsagn DKK"], x["Indbetalt DKK"]
    if ("Resttilsagn DKK" in afviger and indb["afvigelse"] is not None
            and _ens(rest["afvigelse"], -indb["afvigelse"])):
        skriv("Resttilsagn DKK", "Spejler forskellen i indbetalt")

    # nøgletallene: Aleta regner dem af sine egne DKK-beløb
    broeker = {"PICC": (("drawdowns",), "commitment", ("Indbetalt DKK",)),
               "DPI": (("distributions",), "drawdowns", ("Udloddet DKK", "Indbetalt DKK")),
               "RVPI": (("marketValue",), "drawdowns", ("NAV DKK", "Indbetalt DKK")),
               "TVPI": (("marketValue", "distributions"), "drawdowns",
                        ("NAV DKK", "Udloddet DKK", "Indbetalt DKK"))}
    for kontrol, (taeller, naevner, beloeb) in broeker.items():
        if kontrol not in afviger or not a.get(naevner):
            continue
        if not _ens(x[kontrol]["aleta"], sum(a[t] for t in taeller) / a[naevner], 0.0005):
            continue
        foelger = [NAVNE[k] for k in beloeb if k in afviger]
        if not dkk:
            skriv(kontrol, f"Aleta regner af sine DKK-beløb (indbetalt og udloddet til "
                           f"transaktionskurserne, NAV og tilsagn til kursen pr. "
                           f"{ALETA_DATO:%d.%m.%Y}); rapporten af beløbene i {valuta}")
        elif foelger:
            skriv(kontrol, "Følger forskellen i " + " og ".join(foelger))

    # IRR: rapportens strømme regnet, som Aleta gør, og forvalterens eget tal
    if "IRR" in afviger:
        dele = []
        if tx and _ens(xirr(_aletas_irr(tx) + [(ALETA_DATO, a["marketValue"])]), a["irr"],
                       0.00005):
            dele.append(f"Aletas IRR er XIRR over Aletas transaktioner i DKK med NAV pr. "
                        f"{ALETA_DATO:%d.%m.%Y} som sidste strøm (regnet efter)")
        if tx_datoer:
            flyt = {t["rapport_dato"]: t["aleta_dato"] for t in tx_datoer}
            dage = sorted(t["dage"] for t in tx_datoer)
            tekst = (f"{_stroemme(len(tx_datoer))} har en anden dato hos Aleta (typisk "
                     f"{abs(dage[len(dage) // 2])} dage {'før' if dage[len(dage) // 2] > 0 else 'efter'}"
                     " rapportens)")
            irr = (xirr([(flyt.get(d, d), b) for d, b in _uden_nav(f)] + [(ALETA_DATO, n["nav"])])
                   if dkk else _irr_dkk(f, a["marketValue"], flyt))
            if irr is not None:
                tekst += (f"; med Aletas datoer og NAV pr. {ALETA_DATO:%d.%m.%Y} giver rapportens "
                          f"strømme {_pct(irr)}")
                tekst += "" if dkk else " i DKK"
            dele.append(tekst)
        if tx_afviger:
            dele.append(f"{_stroemme(len(tx_afviger))} har et andet beløb eller findes kun det "
                        "ene sted (se Indbetalt og Udloddet)")
        if not dkk and not tx_datoer and (irr := _irr_dkk(f, a["marketValue"])) is not None:
            dele.append(f"i DKK giver rapportens strømme {_pct(irr)}")
        if foer_carry and (irr := xirr(_uden_nav(f) + [(f["nav_dato"], ca890)])) is not None:
            dele.append(f"med NAV før carry giver rapportens strømme {_pct(irr)}")
        if udligning and "Aleta regner udligningsrenten" in indb["bemaerkning"]:
            uden = [(d, b) for d, b in f["irr_flows"] if b != -n["udligning"]]
            if (irr := xirr(uden)) is not None:
                dele.append(f"uden udligningsrenten giver rapportens strømme {_pct(irr)}")
        if tx:
            pass
        elif dkk and "Indbetalt DKK" in afviger and not udligning:
            dele.append("resten følger forskellen i indbetalt" if foer_carry
                        else "indbetalt afviger også")
        elif not afviger & {"NAV", "NAV DKK", "Udloddet DKK"} and dkk:
            dele.append(("ellers stemmer NAV, indbetalt og udloddet, så resten" if udligning
                         else "NAV, indbetalt og udloddet stemmer, så forskellen")
                        + " ligger i kapitalstrømmenes datoer eller Aletas beregning")
        if n.get("gp_irr") is not None:
            dele.append(f"forvalterens IRR er {_pct(n['gp_irr'] / 100)}")
        if dele:
            skriv("IRR", "; ".join(dele)[0].upper() + "; ".join(dele)[1:])

    # vintage
    if "Vintage" in afviger:
        kald = [d for d, _, t in f.get("kapitalstroemme", []) if t == "kald"]
        tekst = f"Rapporten bruger fondens vintage fra stamdata ({f['vintage']})"
        if transaktioner:
            tekst = (f"Aletas vintage er fondens startdato i Aleta "
                     f"({transaktioner[0]['vintage']:%d.%m.%Y}); " + tekst[0].lower() + tekst[1:])
        if kald:
            tekst += f"; første kapitalkald {min(kald):%d.%m.%Y}"
        skriv("Vintage", tekst)


# --- afstemningsdokumentet -----------------------------------------------------------------------

GROEN = PatternFill("solid", fgColor="C6EFCE")
GUL = PatternFill("solid", fgColor="FFEB9C")
ROED = PatternFill("solid", fgColor="FFC7CE")
HOVED = PatternFill("solid", fgColor="4F6B5A")
GRUPPE = PatternFill("solid", fgColor="2F4538")
HVID_FED = Font(bold=True, color="FFFFFF")
GRAA = Font(color="A6A6A6")
STREG = Border(top=Side(style="thin", color="808080"))
GRAA_FYLD = PatternFill("solid", fgColor="EDEDED")
MARKERING = {"ok": ("✓", GROEN), "lille afvigelse": ("≈", GUL), "afviger": ("✗", ROED),
             "efter rapportdatoen": ("–", GRAA_FYLD)}
TX_FORKLARING = ("Aletas transaktioner mod rapportens kapitalstrømme pr. fond, type og dato (beløb i "
                 "fondens valuta; samme dags konti lagt sammen). Grøn: samme dato og beløb. Gul: samme "
                 "beløb, anden dato. Rød: andet beløb eller kun det ene sted. Grå: hos Aleta efter "
                 "rapportdatoen, indgår ikke.")
# fanen Transaktioner: (overskrift, nøgle, bredde, talformat)
TX_KOLONNER = [
    ("Fond", "fond", 11, None), ("Type", "type", 10, None), ("Status", "art", 18, None),
    ("Aleta dato", "aleta_dato", 11, "dd.mm.yyyy"),
    ("Rapport dato", "rapport_dato", 11, "dd.mm.yyyy"), ("Dage", "dage", 6, "0"),
    ("Valuta", "valuta", 7, None), ("Aleta", "aleta", 14, "#,##0.00"),
    ("Rapport", "rapport", 14, "#,##0.00"), ("Afvigelse", "afvigelse", 12, "#,##0.00"),
    ("Aleta DKK", "aleta_dkk", 14, "#,##0"), ("Rapport DKK", "rapport_dkk", 14, "#,##0"),
    ("Afvigelse DKK", "afvigelse_dkk", 12, "#,##0"), ("Aletas kurs", "aleta_kurs", 9, "0.0000"),
    ("Nationalbanken", "rapport_kurs", 9, "0.0000"), ("Bemærkning", "bemaerkning", 60, None),
    ("Rapportens kapitalstrøm", "ref", 45, None), ("Aletas transaktions-ID", "aleta_id", 28, None),
]
FORKLARING = ("Grøn: ens efter afrunding, eller under 1 % af den tilladte afvigelse. Gul: afviger, "
              "men inden for den tilladte afvigelse. "
              "Rød: uden for den tilladte afvigelse.")


def _format(kontrol):
    if kontrol in ("PICC", "IRR"):
        return "0.00%"
    if kontrol in ("DPI", "RVPI", "TVPI"):
        return "0.000"
    return "#,##0"


def _andel(x):
    """Afvigelsen som andel af den tilladte afvigelse."""
    if x["afvigelse"] is None or not x["tilladt"]:
        return None
    return abs(x["afvigelse"]) / x["tilladt"]


# fanen Detaljer: (gruppe, [(overskrift, nøgle, bredde), ...]); nøglen "andel" beregnes af _andel
DETALJER = [
    ("", [("Fond", "fond", 11), ("Kontrol", "kontrol", 16)]),
    ("Sammenligning (beløb i DKK, undtagen NAV)", [
        ("Status", "status", 15), ("Aleta", "aleta", 15), ("Rapport", "rapport", 15),
        ("Afvigelse", "afvigelse", 13), ("Tilladt afvigelse", "tilladt", 14),
        ("Andel af tilladt", "andel", 10)]),
    ("I fondens valuta", [
        ("Valuta", "valuta", 8), ("Aleta", "aleta_fv", 15), ("Rapport", "rapport_fv", 15),
        ("Afvigelse", "afvigelse_fv", 13)]),
    ("Omregningskurs", [("Kurs", "kurs", 9), ("Kursdato", "kursdato", 18)]),
    ("", [("Bemærkning", "bemaerkning", 90), ("Navn i Aleta", "aleta_navn", 34)]),
]


def _fondsstatus(tx, fond):
    """Den værste status blandt fondens transaktioner (strømme efter rapportdatoen tæller ikke)."""
    status = {t["status"] for t in tx if t["fond"] == fond} - {"efter rapportdatoen"}
    return next((s for s in ("afviger", "lille afvigelse", "ok") if s in status), None)


def skriv(linjer, as_of=AS_OF, transaktioner=None):
    """Projektmappe med fanerne Oversigt (fond × kontrol), Detaljer (én række pr. kontrol) og,
    med transaktioner fra afstem_transaktioner(), Transaktioner (én række pr. kapitalstrøm) og IRR
    (Aletas og rapportens betalinger og slutværdi side om side med en XIRR for hver)."""
    wb = Workbook()
    ov = wb.active
    ov.title = "Oversigt"
    antal = {s: sum(x["status"] == s for x in linjer) for s in MARKERING}
    ov["A1"] = f"Motortramp: rapporten mod Aleta pr. {as_of:%d.%m.%Y}"
    ov["A1"].font = Font(bold=True, size=14)
    ov["A2"] = (f"{antal['ok']} ok, {antal['lille afvigelse']} med lille afvigelse, "
                f"{antal['afviger']} afviger - se fanen Detaljer. {FORKLARING}")
    ov["A3"] = (f"Aletas tal er pr. {aleta_dato(as_of):%d.%m.%Y}. Rapportens NAV, tilsagn og "
                "resttilsagn i EUR og USD er omregnet til DKK med Nationalbankens kurs på Aletas "
                "dato, indbetalt og udloddet med kursen på hver transaktionsdato. I Detaljer står "
                "beløbene også i fondens valuta; Aletas DKK-beløb er regnet tilbage med samme kurs.")
    if transaktioner:
        tx_antal = {a: sum(t["art"] == a for t in transaktioner) for a in TRANSAKTIONSSTATUS}
        ov["A4"] = ("Transaktioner: " + ", ".join(f"{v} {a}" for a, v in tx_antal.items() if v)
                    + " - se fanen Transaktioner.")
    kontroller = [k[0] for k in KONTROLLER] + (["Transaktioner"] if transaktioner else [])
    fonde = list(dict.fromkeys(x["fond"] for x in linjer))
    for j, navn in enumerate(["Fond", *kontroller], start=1):
        c = ov.cell(5, j, navn)
        c.fill, c.font, c.alignment = HOVED, HVID_FED, Alignment(horizontal="center")
    opslag = {(x["fond"], x["kontrol"]): x for x in linjer}
    for i, fond in enumerate(fonde, start=6):
        ov.cell(i, 1, fond).font = Font(bold=True)
        for j, kontrol in enumerate(kontroller, start=2):
            status = (_fondsstatus(transaktioner, fond) if kontrol == "Transaktioner"
                      else opslag[fond, kontrol]["status"])
            if status is None:
                continue
            tegn, fyld = MARKERING[status]
            c = ov.cell(i, j, tegn)
            c.fill, c.alignment = fyld, Alignment(horizontal="center")
    ov.column_dimensions["A"].width = 12
    for j in range(2, len(kontroller) + 2):
        ov.column_dimensions[ov.cell(5, j).column_letter].width = 13
    ov.freeze_panes = "B6"

    de = wb.create_sheet("Detaljer")
    de["A1"] = FORKLARING
    de["A1"].font = Font(italic=True)
    kolonner = [k for _, gruppe in DETALJER for k in gruppe]
    sammenligning = {k[1] for k in DETALJER[1][1]}
    j = 1
    for gruppe, felter in DETALJER:
        for _ in felter:
            de.cell(2, j).fill = GRUPPE
            j += 1
        if gruppe:
            c = de.cell(2, j - len(felter), gruppe)
            c.font, c.alignment = HVID_FED, Alignment(horizontal="center")
            de.merge_cells(start_row=2, start_column=j - len(felter), end_row=2, end_column=j - 1)
    for j, (navn, _, bredde) in enumerate(kolonner, start=1):
        c = de.cell(3, j, navn)
        c.fill, c.font = HOVED, HVID_FED
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        de.column_dimensions[c.column_letter].width = bredde
    forrige = None
    for i, x in enumerate(linjer, start=4):
        ny_fond = x["fond"] != forrige
        forrige = x["fond"]
        eksakt = x["kontrol"] in ("Valuta", "Vintage")
        for j, (_, noegle, _) in enumerate(kolonner, start=1):
            c = de.cell(i, j, _andel(x) if noegle == "andel" else x[noegle])
            if noegle in ("aleta", "rapport", "afvigelse", "tilladt") and not eksakt:
                c.number_format = _format(x["kontrol"])
            elif noegle in ("aleta_fv", "rapport_fv", "afvigelse_fv"):
                c.number_format = "#,##0"
            elif noegle == "andel":
                c.number_format = "0%"
            elif noegle == "kurs":
                c.number_format = "0.0000"
            elif noegle == "kursdato":
                c.number_format = "dd.mm.yyyy"
            if noegle in sammenligning:
                c.fill = MARKERING[x["status"]][1]
            if noegle == "fond":
                c.font = Font(bold=True) if ny_fond else GRAA
            if ny_fond and i > 4:
                c.border = STREG
    de.auto_filter.ref = f"A3:{de.cell(3, len(kolonner)).column_letter}{len(linjer) + 3}"
    de.freeze_panes = "C4"
    if transaktioner:
        _skriv_transaktioner(wb.create_sheet("Transaktioner"), transaktioner)
        _skriv_irr(wb.create_sheet("IRR"), linjer, transaktioner, aleta_dato(as_of))
    return wb


def _skriv_transaktioner(ws, transaktioner):
    ws["A1"] = TX_FORKLARING
    ws["A1"].font = Font(italic=True)
    for j, (navn, _, bredde, _) in enumerate(TX_KOLONNER, start=1):
        c = ws.cell(2, j, navn)
        c.fill, c.font = HOVED, HVID_FED
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = bredde
    forrige = None
    for i, t in enumerate(transaktioner, start=3):
        ny_fond = t["fond"] != forrige
        forrige = t["fond"]
        for j, (_, noegle, _, fmt) in enumerate(TX_KOLONNER, start=1):
            c = ws.cell(i, j, t[noegle])
            if fmt:
                c.number_format = fmt
            if noegle in ("art", "aleta_dato", "rapport_dato", "aleta", "rapport", "afvigelse"):
                c.fill = MARKERING[t["status"]][1]
            if noegle == "fond":
                c.font = Font(bold=True) if ny_fond else GRAA
            if ny_fond and i > 3:
                c.border = STREG
    ws.auto_filter.ref = f"A2:{ws.cell(2, len(TX_KOLONNER)).column_letter}{len(transaktioner) + 2}"
    ws.freeze_panes = "D3"


IRR_FORKLARING = ("Aletas IRR genberegnet som XIRR over Aletas transaktioner i DKK med markedsværdien "
                  "på Aletas dato som slutværdi, ved siden af rapportens IRR i DKK: rapportens "
                  "betalinger med Nationalbankens kurs på hver betalingsdag og NAV med kursen på "
                  "Aletas dato. Betalingerne er parret som på fanen Transaktioner; mangler en "
                  "betaling på den ene side, står den som 0 (grå), så den ikke påvirker XIRR. "
                  "Strømme efter rapportdatoen er udeladt, som i begge IRR. For DKK-fonde er "
                  "decket IRR regnet med NAV pr. rapportdatoen, her pr. Aletas dato, så de to "
                  "rapport-IRR kan afvige en smule. XIRR-rækken er formler, så tal kan rettes og "
                  "IRR regnes om.")
IRR_KOLONNER = [("Type", 11), ("Status", 18), ("Aleta dato", 11), ("Aleta DKK", 15),
                ("Rapport dato", 11), ("Rapport DKK", 15), ("Dage", 6), ("Forskel DKK", 13),
                ("Bemærkning", 60)]


def _skriv_irr(ws, linjer, transaktioner, dato):
    """Fanen IRR: én blok pr. fond med de parrede betalinger (kald negative), slutværdien og en
    XIRR-formel for Aleta og for rapporten, Aletas egen IRR og rapportens IRR i DKK fra
    kontrollerne, og forskellen."""
    ws["A1"] = IRR_FORKLARING
    ws["A1"].font = Font(italic=True)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 60
    ws.column_dimensions["A"].width = 11
    for j, (_, bredde) in enumerate(IRR_KOLONNER, start=2):
        ws.column_dimensions[get_column_letter(j)].width = bredde
    opslag = {(x["fond"], x["kontrol"]): x for x in linjer}
    i = 3
    for fond in dict.fromkeys(x["fond"] for x in linjer):
        tx = [t for t in transaktioner if t["fond"] == fond and t["art"] != "efter rapportdatoen"]
        nav, irr = opslag.get((fond, "NAV DKK")), opslag.get((fond, "IRR"))
        if not tx or nav is None:
            continue
        c = ws.cell(i, 1, fond)
        c.font = Font(bold=True, size=12)
        i += 1
        for j, (navn, _) in enumerate(IRR_KOLONNER, start=2):
            c = ws.cell(i, j, navn)
            c.fill, c.font = HOVED, HVID_FED
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        i += 1
        foerste = i
        for t in sorted(tx, key=lambda t: t["aleta_dato"] or t["rapport_dato"]):
            fortegn = -1 if t["type"] == "kald" else 1
            a_dato = t["aleta_dato"] or t["rapport_dato"]
            r_dato = t["rapport_dato"] or t["aleta_dato"]
            a_dkk = fortegn * t["aleta_dkk"] if t["aleta_dkk"] is not None else 0
            r_dkk = fortegn * t["rapport_dkk"] if t["rapport_dkk"] is not None else 0
            vaerdier = [t["type"], t["art"], a_dato, a_dkk, r_dato, r_dkk, t["dage"],
                        f"=G{i}-E{i}", t["bemaerkning"]]
            for j, v in enumerate(vaerdier, start=2):
                c = ws.cell(i, j, v)
                c.number_format = {4: "dd.mm.yyyy", 6: "dd.mm.yyyy", 5: "#,##0", 7: "#,##0",
                                   8: "0", 9: "#,##0"}.get(j, "General")
            ws.cell(i, 3).fill = MARKERING[t["status"]][1]
            # en betaling, der kun findes på den ene side, står som 0 på den anden
            if t["aleta_dkk"] is None:
                ws.cell(i, 4).font = ws.cell(i, 5).font = GRAA
            if t["rapport_dkk"] is None:
                ws.cell(i, 6).font = ws.cell(i, 7).font = GRAA
            i += 1
        # slutværdien: Aletas markedsværdi og rapportens NAV omregnet på Aletas dato
        vaerdier = ["slutværdi", "NAV", dato, nav["aleta"], dato, nav["rapport"], None,
                    f"=G{i}-E{i}",
                    f"rapportens NAV {nav['rapport_fv']:,.0f} {nav['valuta']} × kurs "
                    f"{nav['kurs'] or 1:.4f}".replace(",", ".") if nav["valuta"] != "DKK"
                    else "markedsværdi / NAV"]
        for j, v in enumerate(vaerdier, start=2):
            c = ws.cell(i, j, v)
            c.number_format = {4: "dd.mm.yyyy", 6: "dd.mm.yyyy", 5: "#,##0", 7: "#,##0",
                               9: "#,##0"}.get(j, "General")
            c.border = STREG
            c.font = Font(bold=True)
        sidste = i
        i += 1
        rammer = [("IRR genberegnet (XIRR)", f"=XIRR(E{foerste}:E{sidste},D{foerste}:D{sidste})",
                   f"=XIRR(G{foerste}:G{sidste},F{foerste}:F{sidste})"),
                  ("IRR i nøgletallene / decket", irr and irr["aleta"], irr and irr["rapport"])]
        for tekst, a, r in rammer:
            ws.cell(i, 2, tekst).font = Font(bold=True)
            for j, v in ((5, a), (7, r)):
                c = ws.cell(i, j, v)
                c.number_format, c.font = "0.0000%", Font(bold=True)
            ws.cell(i, 9, f"=(G{i}-E{i})*10000").number_format = '0.0" bp"'
            i += 1
        ws.cell(i - 2, 10, "rapporten - Aleta i basispoint")
        ws.cell(i - 1, 10, "genberegnet mod nøgletal: Aleta E" + str(i - 2) + " vs E"
                + str(i - 1) + ", rapporten G" + str(i - 2) + " vs G" + str(i - 1))
        i += 2
    ws.freeze_panes = "A3"
