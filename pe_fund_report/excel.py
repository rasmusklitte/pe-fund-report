"""Dataudtræk til Excel: tallene bag rapporten med kildedokument på hver række.

Filen er til intern afstemning, ikke til klienten: den viser dokumentstier og interne
id'er. Hver række, der kommer fra databasen, har sit dokument-id (link til fanen Dokumenter)
og et link til PDF'en. Afledte tal (PICC, DPI, TVPI, resttilsagn, værdiændring, DKK-beløb)
står som Excel-formler, og kontrolkolonnerne summerer kapitalstrømmene og regner IRR med
XIRR, så hvert nøgletal kan følges tilbage til de enkelte dokumenter.

Som slides.py læser modulet kun dicten fra data.investor_rapport(); det spørger ikke databasen.
"""
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from . import style
from .paths import DOKUMENTER

FORMATER = {"tal": "#,##0;-#,##0;0", "dec": "#,##0.00", "pct": "0.0%", "gange": '0.00"x"',
            "dato": "dd-mm-yyyy", "kurs": "0.0000", "tekst": "@", None: "General"}
HOVED = PatternFill("solid", fgColor=style.PPIM_SAGE_DARK.lstrip("#"))
KONTROL = PatternFill("solid", fgColor=style.PPIM_TERRACOTTA.lstrip("#"))
KILDE = PatternFill("solid", fgColor=style.PPIM_TAUPE_LIGHT.lstrip("#"))
HVID_FED = Font(bold=True, color="FFFFFF")
LINK = Font(color="0563C1", underline="single")

PERIODER = {"kvartal": "Kvartalet", "aar": "År til dato"}
KS = "Kapitalstrømme"       # fanenavne, der bruges i formler på tværs af faner


class Link:
    """Celletekst med et hyperlink: en fil på disken eller en celle i projektmappen ("#Ark!A1")."""

    def __init__(self, tekst, maal):
        self.tekst, self.maal = tekst, maal


def _v(v):
    """Pandas/numpy-værdier til noget, openpyxl kan skrive; NaN og NaT bliver tomme."""
    if v is None:
        return None
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.date() if not pd.isna(v) else None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v.item() if hasattr(v, "item") else v


class Ark:
    """Én fane: kolonner er [(navn, bredde, format, art)], art None, "kontrol" eller "kilde".

    Rækkerne skrives fra række 2. En celleværdi kan være en funktion af rækkenummeret og en
    {kolonnenavn: bogstav}-ordbog, som returnerer en formel - så formlerne følger kolonnerne,
    hvis rækkefølgen ændres.
    """

    def __init__(self, wb, navn, kolonner):
        self.ws = wb.create_sheet(navn)
        self.navn, self.kolonner = navn, kolonner
        self.k = {n: get_column_letter(i) for i, (n, *_) in enumerate(kolonner, 1)}
        self.rad = 1
        for i, (n, bredde, _, art) in enumerate(kolonner, 1):
            c = self.ws.cell(1, i, n)
            c.font, c.fill = HVID_FED, {"kontrol": KONTROL, "kilde": KILDE}.get(art, HOVED)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            self.ws.column_dimensions[get_column_letter(i)].width = bredde
        self.ws.row_dimensions[1].height = 30
        self.ws.freeze_panes = "C2"

    def kol(self, navn):
        """Hele kolonnen som område til SUMIFS fra en anden fane: 'Ark'!$H:$H."""
        return f"'{self.navn}'!${self.k[navn]}:${self.k[navn]}"

    def raekke(self, vaerdier, fed=False):
        self.rad += 1
        for i, ((_, _, fmt, _), v) in enumerate(zip(self.kolonner, vaerdier), 1):
            if callable(v):
                v = v(self.rad, self.k)
            c = self.ws.cell(self.rad, i)
            if isinstance(v, Link):
                c.value, c.font = v.tekst, LINK
                # "#Ark!A1" er et hop i projektmappen, alt andet en fil
                c.hyperlink = (Hyperlink(ref=c.coordinate, location=v.maal[1:])
                               if v.maal.startswith("#") else v.maal)
            else:
                c.value = _v(v)
            c.number_format = FORMATER[fmt]
            if fed:
                c.font = Font(bold=True)
        return self.rad

    def afslut(self):
        if self.rad > 1:
            self.ws.auto_filter.ref = f"A1:{get_column_letter(len(self.kolonner))}{self.rad}"


class Kilder:
    """Opslag fra dokument-id til fil og til rækken i fanen Dokumenter."""

    def __init__(self, dokumenter):
        self.dok = {int(r.document_id): r for r in dokumenter.itertuples()}
        self.rad = {}           # dokument-id -> række i fanen Dokumenter, sat når den skrives

    def id_(self, doc_id):
        if doc_id is None or pd.isna(doc_id):
            return None
        doc_id = int(doc_id)
        return Link(doc_id, f"#Dokumenter!A{self.rad[doc_id]}") if doc_id in self.rad else doc_id

    def fil(self, doc_id):
        if doc_id is None or pd.isna(doc_id) or int(doc_id) not in self.dok:
            return None
        d = self.dok[int(doc_id)]
        sti = DOKUMENTER / d.rel_path if d.rel_path else None
        return Link(d.file_name, sti.as_uri()) if sti and sti.exists() else d.file_name


def _fond_navn(r):
    return {f["fund_code"]: f["fund_name"] for f in r["fonde"]}


# ---------- fanerne ----------

def _dokumenter(wb, r, kilder):
    ark = Ark(wb, "Dokumenter", [
        ("Dokument-id", 11, None, None), ("Fondskode", 11, None, None), ("Type", 18, None, None),
        ("Periode", 11, "dato", None), ("Valørdato", 11, "dato", None),
        ("Reference", 24, None, None), ("Fil", 50, None, "kilde"), ("Sider", 7, None, None),
        ("Status", 10, None, None), ("Udtræk", 10, None, None), ("Godkendt af", 17, None, None),
        ("Godkendt", 17, "dato", None), ("Pladsholder", 11, None, "kontrol"),
        ("Note", 60, None, None)])
    for d in r["dokumenter"].itertuples():
        pladsholder = "Ja" if "PLACEHOLDER" in (d.file_name or "").upper() else ""
        kilder.rad[int(d.document_id)] = ark.raekke([
            int(d.document_id), d.fund_code, d.doc_type, d.period_end, d.value_date,
            d.external_ref, kilder.fil(d.document_id), d.page_count, d.status,
            d.extraction_method, d.approved_by, d.approved_at, pladsholder, d.note])
    ark.afslut()
    return ark


def _kapitalstroemme(wb, r, kilder):
    ark = Ark(wb, KS, [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Valørdato", 11, "dato", None),
        ("Meddelelsesdato", 11, "dato", None), ("Konto", 8, None, None),
        ("Kontonavn", 30, None, None), ("Klasse", 16, None, None),
        ("Genkaldelig", 10, None, None), ("Beløb", 14, "tal", None), ("Valuta", 7, None, None),
        ("Kurs", 9, "kurs", None), ("Beløb DKK", 14, "tal", None),
        ("I nøgletal", 9, None, None), ("I IRR", 7, None, None), ("Bemærkning", 44, None, None),
        ("Reference", 30, None, None), ("Dokument-id", 11, None, "kilde"),
        ("Kildedokument", 50, None, "kilde")])
    for f in r["fonde"]:
        for c in f["kilder"]["cash_flows"].itertuples():
            efter = c.value_date > f["nav_dato"]
            bem = ("Efter NAV-datoen – ikke medregnet" if efter else
                   {"EQUALISATION": "Udligningsrente – i indbetalt (betalt) eller udloddet "
                                    "(modtaget), ikke i kaldt kapital",
                    "WITHHOLDING_TAX": "Kildeskat – holdes ude af IRR og DPI"}.get(c.flow_class, ""))
            dkk = c.amount_dkk if not pd.isna(c.amount_dkk) else (
                c.amount if f["valuta"] == "DKK" else None)
            ark.raekke([
                f["fund_code"], f["fund_name"], c.value_date, c.notice_date, c.account_code,
                c.account_name, c.flow_class, "Ja" if c.is_recallable == 1 else "",
                c.amount, f["valuta"], c.fx_rate if not pd.isna(c.fx_rate) else
                (1.0 if f["valuta"] == "DKK" else None), dkk,
                "Ja" if not efter and c.flow_class in ("CONTRIBUTION", "DISTRIBUTION",
                                                       "EQUALISATION") else "Nej",
                "Ja" if not efter and c.flow_class != "WITHHOLDING_TAX" else "Nej",
                bem, c.flow_ref, kilder.id_(c.document_id), kilder.fil(c.document_id)])
        for e in f["kilder"]["estimerede"]:
            ark.raekke([
                f["fund_code"], f["fund_name"], pd.Timestamp(e["dato"]).date(), None, None,
                "Skønnet kapitalkald", "CONTRIBUTION", "", float(e["beloeb"]), f["valuta"],
                None, None, "Nej", "Ja",
                f"Skønnet – meddelelsen er ikke modtaget: {e.get('note', '')}", None, None,
                "config/fund_info.yaml"])
    ark.afslut()
    return ark


def _nav(wb, r, kilder):
    ark = Ark(wb, "NAV", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Periode", 11, "dato", None),
        ("NAV rapporteret", 15, "tal", None), ("Restatement", 13, "tal", None),
        ("NAV brugt", 15, "tal", None), ("Brugt som", 26, None, None),
        ("Dokument-id", 11, None, "kilde"), ("Kildedokument", 50, None, "kilde")])
    nav_rad = {}           # (fondskode, periode) -> række, til links fra Nøgletal og Udvikling
    for f in r["fonde"]:
        roller = {f["nav_dato"]: "NAV ultimo (rapportens NAV)"}
        for p, u in f["udvikling"].items():
            roller.setdefault(u["primo_dato"], f"NAV primo, {PERIODER[p].lower()}")
        for n in f["kilder"]["nav"].itertuples():
            nav_rad[f["fund_code"], n.period_end] = ark.raekke([
                f["fund_code"], f["fund_name"], n.period_end, n.nav_rapporteret,
                n.restatement, n.nav, roller.get(n.period_end, ""), kilder.id_(n.document_id),
                kilder.fil(n.document_id)])
    ark.afslut()
    return nav_rad


BRUGT_TIL = {"CA110": "Forvalterens indbetalte kapital (ITD)", "MM110": "Kaldt mod tilsagnet (resttilsagn; forvalterens indbetalte, hvis CA110 mangler)",
             "CA210": "Værdibro: gebyrer", "CA220": "Værdibro: omkostninger",
             "CA230": "Værdibro: finansielle poster", "CA400": "Værdibro: carried interest",
             "MM150": "Forvalterens IRR", "MM160": "Forvalterens TVPI"}


def _kapitalkonto(wb, r, kilder):
    ark = Ark(wb, "Kapitalkonto", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Periode", 11, "dato", None),
        ("Basis", 7, None, None), ("Konto", 8, None, None), ("Kontonavn", 36, None, None),
        ("Beløb", 15, "dec", None), ("Brugt til", 34, None, None),
        ("Dokument-id", 11, None, "kilde"), ("Kildedokument", 50, None, "kilde")])
    for f in r["fonde"]:
        ka = f["kilder"]["kapitalkonto"].sort_values(["period_end", "basis", "account_code"])
        for k in ka.itertuples():
            brugt = BRUGT_TIL.get(k.account_code, "") if k.period_end == f["nav_dato"] else ""
            ark.raekke([f["fund_code"], f["fund_name"], k.period_end, k.basis, k.account_code,
                        k.account_name, k.amount, brugt, kilder.id_(k.document_id),
                        kilder.fil(k.document_id)])
    ark.afslut()


def _tilsagn(wb, r, kilder):
    ark = Ark(wb, "Tilsagn", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Dato", 11, "dato", None),
        ("Type", 12, None, None), ("Beløb", 15, "tal", None), ("Note", 60, None, None),
        ("Dokument-id", 11, None, "kilde"), ("Kildedokument", 50, None, "kilde")])
    for f in r["fonde"]:
        for t in f["kilder"]["tilsagn"].itertuples():
            ark.raekke([f["fund_code"], f["fund_name"], t.event_date, t.event_type, t.amount,
                        t.note, kilder.id_(t.document_id), kilder.fil(t.document_id)])
    ark.afslut()
    return ark


def _irr(wb, r):
    """Hver fonds IRR-cash flows i én blok, sorteret efter dato, så XIRR kan regne på dem."""
    ark = Ark(wb, "IRR", [("Fondskode", 11, None, None), ("Fond", 34, None, None),
                          ("Dato", 11, "dato", None), ("Beløb", 15, "tal", None),
                          ("Type", 24, None, None)])
    blokke = {}
    for f in r["fonde"]:
        flows = f["irr_flows"]
        antal_est = len(f["kilder"]["estimerede"])
        typer = (["Cash flow"] * (len(flows) - antal_est - 1)
                 + ["Skønnet kapitalkald"] * antal_est + ["NAV (sidste cash flow)"])
        raekker = sorted(zip(flows, typer), key=lambda x: (x[0][0], x[1] == "NAV (sidste cash flow)"))
        foerste = ark.rad + 1
        for (d, a), t in raekker:
            ark.raekke([f["fund_code"], f["fund_name"], d, a, t])
        blokke[f["fund_code"]] = (foerste, ark.rad)
    ark.afslut()
    return blokke


def _noegletal(wb, r, kilder, ks, nav_rad, irr_blokke):
    ark = Ark(wb, "Nøgletal", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Valuta", 7, None, None),
        ("NAV-dato", 11, "dato", None), ("Tilsagn", 14, "tal", None),
        ("Indbetalt", 14, "tal", None), ("Heraf udligningsrente", 12, "tal", None),
        ("Udloddet", 14, "tal", None), ("Kaldt mod tilsagn", 14, "tal", None), ("Rabat på kald", 11, "tal", None),
        ("Genkaldelige udlodninger", 13, "tal", None), ("NAV", 14, "tal", None),
        ("Resttilsagn", 14, "tal", None), ("PICC", 8, "pct", None), ("DPI", 8, "gange", None),
        ("TVPI", 8, "gange", None), ("IRR", 8, "pct", None),
        ("Fondsstørrelse", 16, "tal", None), ("Ejerandel", 9, "pct", None),
        ("DKK-kurs", 9, "kurs", None), ("NAV DKK", 14, "tal", None),
        ("Kontrol: indbetalt iflg. kapitalkald", 15, "tal", "kontrol"),
        ("Kontrol: afvigelse", 12, "tal", "kontrol"),
        ("Kontrol: udloddet iflg. udlodninger", 15, "tal", "kontrol"),
        ("Kontrol: IRR med XIRR", 10, "pct", "kontrol"),
        ("Kilde tilsagn", 22, None, "kilde"), ("Kilde indbetalt", 34, None, "kilde"),
        ("Kilde NAV", 50, None, "kilde"), ("Kilde fondsstørrelse", 40, None, "kilde"),
        ("Noter", 80, None, None)])
    # udligningsrenten tæller med: betalt (< 0) i indbetalt, modtaget (> 0) i udloddet
    sumifs = lambda klasse, fortegn: (
        lambda rad, k: f'(SUMIFS({ks.kol("Beløb")},{ks.kol("Fondskode")},{k["Fondskode"]}{rad},'
                       f'{ks.kol("Klasse")},"{klasse}",{ks.kol("I nøgletal")},"Ja")'
                       f'+SUMIFS({ks.kol("Beløb")},{ks.kol("Fondskode")},{k["Fondskode"]}{rad},'
                       f'{ks.kol("Klasse")},"EQUALISATION",{ks.kol("I nøgletal")},"Ja",'
                       f'{ks.kol("Beløb")},"{fortegn}0"))')
    for f in r["fonde"]:
        n, kl = f["n"], f["kilder"]
        kurs = (r["kurser"].get(f["valuta"]) or (None,))[0]
        nav_doc = kl["nav"].document_id.iloc[-1]
        fs = kl["fondsstoerrelse"]
        foerste, sidste = irr_blokke[f["fund_code"]]
        ark.raekke([
            f["fund_code"], f["fund_name"], f["valuta"], f["nav_dato"], n["tilsagn"],
            n["indbetalt"], n.get("udligning", 0.0), n["udloddet"],
            n.get("kaldt", n["indbetalt"]),
            lambda rad, k: f"={k['Kaldt mod tilsagn']}{rad}-{k['Indbetalt']}{rad}"
                           f"+{k['Heraf udligningsrente']}{rad}",
            n["recallable"], n["nav"],
            lambda rad, k: f"={k['Tilsagn']}{rad}-{k['Kaldt mod tilsagn']}{rad}"
                           f"+{k['Genkaldelige udlodninger']}{rad}",
            lambda rad, k: f"=IFERROR({k['Indbetalt']}{rad}/{k['Tilsagn']}{rad},\"\")",
            lambda rad, k: f"=IFERROR({k['Udloddet']}{rad}/{k['Indbetalt']}{rad},\"\")",
            lambda rad, k: f"=IFERROR(({k['Udloddet']}{rad}+{k['NAV']}{rad})"
                           f"/{k['Indbetalt']}{rad},\"\")",
            n["irr"], n["fondsstoerrelse"],
            # evergreen: andelen af fondens NAV (fondssiden siger det samme i noten)
            lambda rad, k, fn=n.get("fonds_nav"): (
                f"=IFERROR({k['NAV']}{rad}/{fn},\"\")" if fn
                else f"=IFERROR({k['Tilsagn']}{rad}/{k['Fondsstørrelse']}{rad},\"\")"),
            kurs, lambda rad, k: f"=IFERROR({k['NAV']}{rad}*{k['DKK-kurs']}{rad},\"\")",
            lambda rad, k: "=-" + sumifs("CONTRIBUTION", "<")(rad, k),
            lambda rad, k: f"={k['Indbetalt']}{rad}-{k['Kontrol: indbetalt iflg. kapitalkald']}{rad}",
            lambda rad, k: "=" + sumifs("DISTRIBUTION", ">")(rad, k),
            f"=IFERROR(XIRR(IRR!$D${foerste}:$D${sidste},IRR!$C${foerste}:$C${sidste}),\"\")",
            Link("Fanen Tilsagn", "#Tilsagn!A1"),
            ("Forvalterens opgørelse (fanen Kapitalkonto, CA110/MM110)"
             if kl["indbetalt"] == "forvalteren" else "Sum af kapitalkald (fanen Kapitalstrømme)"),
            kilder.fil(nav_doc) or Link("Fanen NAV", f"#NAV!A{nav_rad[f['fund_code'], f['nav_dato']]}"),
            (f"{fs['tekst']}: {kilder.dok[fs['document_id']].file_name}"
             if fs["document_id"] in kilder.dok else fs["tekst"]),
            " ".join(f["noter"])])
    ark.afslut()


def _vaerdibro(wb, r):
    ark = Ark(wb, "Værdibro", [("Fondskode", 11, None, None), ("Fond", 34, None, None),
                               ("Post", 42, None, None), ("Beløb", 15, "tal", None),
                               ("Kilde", 60, None, "kilde")])
    kilde = {"Indbetalt kapital": "Nøgletal: indbetalt", "Udlodninger": "Nøgletal: udloddet",
             "Forvaltningsgebyr og omkostninger": "Kapitalkonto CA210 + CA220 + CA230, ITD",
             "Carried interest": "Kapitalkonto CA400 (ITD, eller summen af YTD ved årsultimo)",
             "Værdiskabelse efter omkostninger og carry":
                 "Residual: NAV - indbetalt + udloddet (forvalteren opgør resultatet samlet, CA290)",
             "NAV": "Nøgletal: NAV"}
    for f in r["fonde"]:
        foerste = ark.rad + 1
        for post, beloeb in f["bro"]:
            ark.raekke([f["fund_code"], f["fund_name"], post, beloeb,
                        kilde.get(post, "Residual: NAV - indbetalt + udloddet - gebyrer - carry")])
        ark.raekke([f["fund_code"], f["fund_name"], "Kontrol: summen af posterne før NAV - NAV",
                    lambda rad, k, a=foerste: f"=SUM({k['Beløb']}{a}:{k['Beløb']}{rad - 2})"
                                              f"-{k['Beløb']}{rad - 1}",
                    "Skal være 0"], fed=True)
    ark.afslut()


def _udvikling(wb, r, ks, nav_rad):
    ark = Ark(wb, "Udvikling", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Periode", 12, None, None),
        ("Primodato", 11, "dato", None), ("Ultimodato", 11, "dato", None),
        ("Valuta", 7, None, None), ("NAV primo", 14, "tal", None),
        ("Kapitalkald", 14, "tal", None), ("Udlodninger", 14, "tal", None),
        ("NAV ultimo", 14, "tal", None), ("Værdiændring", 14, "tal", None),
        ("Afkast (Modified Dietz)", 11, "pct", None),
        ("DKK-kurs primo", 10, "kurs", None), ("DKK-kurs ultimo", 10, "kurs", None),
        ("NAV primo DKK", 14, "tal", None), ("Kapitalkald DKK", 14, "tal", None),
        ("Udlodninger DKK", 14, "tal", None), ("Værdiændring DKK", 14, "tal", None),
        ("Valutakursændring DKK", 14, "tal", None), ("NAV ultimo DKK", 14, "tal", None),
        ("Kontrol: kapitalkald iflg. kapitalstrømme", 15, "tal", "kontrol"),
        ("Kontrol: udlodninger iflg. kapitalstrømme", 15, "tal", "kontrol"),
        ("Kilde NAV primo", 14, None, "kilde"), ("Kilde NAV ultimo", 14, None, "kilde")])
    kursopslag = r["kursopslag"]

    def kurs(valuta, d):
        return (kursopslag.get((valuta, d)) or (None,))[0]

    def flow_sum(klasse, fortegn):
        # plus udligningsrenten: betalt (< 0) hører til kaldene, modtaget (> 0) til udlodningerne
        beloeb = ks.kol("Beløb")
        ekstra = f',{beloeb},"{"<" if fortegn == "-" else ">"}0"'

        def summ(kl, rad, k, kriterie=""):
            return (f'SUMIFS({beloeb},{ks.kol("Fondskode")},{k["Fondskode"]}{rad},'
                    f'{ks.kol("Klasse")},"{kl}",{ks.kol("Valørdato")},">"&{k["Primodato"]}{rad},'
                    f'{ks.kol("Valørdato")},"<="&{k["Ultimodato"]}{rad}{kriterie})')
        return lambda rad, k: (f"={fortegn}({summ(klasse, rad, k)}"
                               f"+{summ('EQUALISATION', rad, k, ekstra)})")

    def nav_link(kode, d):
        rad = nav_rad.get((kode, d))
        return Link(f"NAV pr. {d:%d.%m.%Y}", f"#NAV!A{rad}") if rad else None

    for p in r["perioder"]:
        for f in r["fonde"]:
            u = f["udvikling"][p]
            ark.raekke([
                f["fund_code"], f["fund_name"], PERIODER[p], u["primo_dato"], f["nav_dato"],
                f["valuta"], u["primo"], u["kald"], u["udlodning"], u["ultimo"],
                lambda rad, k: f"={k['NAV ultimo']}{rad}-{k['NAV primo']}{rad}"
                               f"-{k['Kapitalkald']}{rad}+{k['Udlodninger']}{rad}",
                u["afkast"], kurs(f["valuta"], u["primo_dato"]), kurs(f["valuta"], u["til"]),
                lambda rad, k: f"={k['NAV primo']}{rad}*{k['DKK-kurs primo']}{rad}",
                u.get("kald_dkk"), u.get("udlodning_dkk"),
                lambda rad, k: f"={k['Værdiændring']}{rad}*{k['DKK-kurs ultimo']}{rad}",
                lambda rad, k: f"={k['NAV ultimo DKK']}{rad}-{k['NAV primo DKK']}{rad}"
                               f"-{k['Kapitalkald DKK']}{rad}+{k['Udlodninger DKK']}{rad}"
                               f"-{k['Værdiændring DKK']}{rad}",
                lambda rad, k: f"={k['NAV ultimo']}{rad}*{k['DKK-kurs ultimo']}{rad}",
                flow_sum("CONTRIBUTION", "-"), flow_sum("DISTRIBUTION", ""),
                nav_link(f["fund_code"], u["primo_dato"]), nav_link(f["fund_code"], f["nav_dato"])])
        if r["udvikling"].get(p):
            # porteføljen: summen af fondsrækkerne for perioden ovenfor
            antal = len(r["fonde"])
            sum_ = lambda kol, n=antal: (lambda rad, k: f"=SUM({k[kol]}{rad - n}:{k[kol]}{rad - 1})")
            ark.raekke(["I alt", "Porteføljen i DKK", PERIODER[p], None, None, "DKK",
                        None, None, None, None, None, r["udvikling"][p]["afkast"], None, None,
                        sum_("NAV primo DKK"), sum_("Kapitalkald DKK"), sum_("Udlodninger DKK"),
                        sum_("Værdiændring DKK"), sum_("Valutakursændring DKK"),
                        sum_("NAV ultimo DKK"), None, None, None, None], fed=True)
    ark.afslut()


def _beholdninger(wb, r, kilder):
    ark = Ark(wb, "Beholdninger", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Rolle", 16, None, None),
        ("Gennem", 10, None, None), ("Rapportdato", 11, "dato", None), ("Navn", 34, None, None),
        ("Type", 12, None, None), ("Instrument", 18, None, None),
        ("Sektor", 22, None, None), ("Branche (GICS)", 18, None, None),
        ("Branche foreløbig", 9, None, None),
        ("Vintage", 8, None, None), ("Closing", 11, "dato", None),
        ("Kostpris, fonden", 15, "tal", None), ("Dagsværdi, fonden", 15, "tal", None),
        ("Udloddet, fonden", 15, "tal", None), ("Ejerandel", 9, "pct", None),
        ("Kostpris, andel", 14, "tal", None), ("Dagsværdi, andel", 14, "tal", None),
        ("Udloddet, andel", 14, "tal", None), ("Værdiskabt, andel", 14, "tal", None),
        ("Skaleret", 8, None, None), ("Dokument-id", 11, None, "kilde"),
        ("Kildedokument", 50, None, "kilde")])
    andel = lambda kol: (lambda rad, k: f"={k[kol + ', fonden']}{rad}*{k['Ejerandel']}{rad}")
    for f in r["fonde"]:
        roller = [("Ultimo", f["beholdninger"])] + [
            (f"Primo, {PERIODER[p].lower()}", b) for p, b in (f.get("beholdninger_primo") or {}).items()]
        for rolle, beh in roller:
            for h in (beh or {}).get("raekker", []):
                # for en feeder er tallene den underliggende fonds, og ejerandelen er NAV-andelen
                gennem = beh.get("kode") if beh.get("kode") != f["fund_code"] else None
                ark.raekke([
                    f["fund_code"], f["fund_name"], rolle, gennem, beh["dato"], h["navn"],
                    h.get("type"), h.get("instrument"),
                    h["sektor"], h.get("industri"),
                    "Ja" if h.get("industri_foreloebig") else None,
                    h["vintage"], h["closing"], h["kostpris_fond"],
                    h["dagsvaerdi_fond"], h["udloddet_fond"], beh["andel"], andel("Kostpris"),
                    andel("Dagsværdi"), andel("Udloddet"),
                    lambda rad, k: f"={k['Dagsværdi, andel']}{rad}+{k['Udloddet, andel']}{rad}"
                                   f"-{k['Kostpris, andel']}{rad}",
                    "Ja" if beh["skaleret"] else "Nej – hele fonden",
                    kilder.id_(h["document_id"]), kilder.fil(h["document_id"])])
    ark.afslut()


GRUNDLAG = {"GP_BREAKDOWN": "Forvalterens fordeling",
            "HOLDINGS": "Beholdningernes dagsværdi", "FUND_DEFAULT": "Fondens strategi"}
# fane, kolonnenavne for de to niveauer og koden under dem, pr. art i data.ARTER
EKSPONERINGSFANER = {"brancher": ("Brancher", "Sektor", "Branchegruppe", "GICS"),
                     "regioner": ("Geografi", "Region", "Delregion", "Kode")}


def _eksponering(wb, r, kilder, art):
    """Siden Branchefordeling eller Geografisk eksponering: hver fonds andele ganget med
    fondens NAV i DKK."""
    if not r.get(art):
        return
    fane, top, under, kode = EKSPONERINGSFANER[art]
    ark = Ark(wb, fane, [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None), ("Gennem", 10, None, None),
        ("Grundlag", 26, None, None), ("Fordelingsdato", 12, "dato", None),
        ("NAV-dato", 11, "dato", None), (top, 26, None, None),
        (under, 36, None, None), (kode, 9, None, None), ("Andel", 9, "pct", None),
        ("Fondens NAV DKK", 15, "tal", None), ("Beløb DKK", 14, "tal", None),
        ("Foreløbig", 9, "pct", None), ("Dokument-id", 11, None, "kilde"),
        ("Kildedokument", 50, None, "kilde")])
    for f in r["fonde"]:
        b = f.get(art)
        linjer = b["linjer"] if b else [dict(top="Ikke klassificeret", under=None,
                                             under_kode=None, andel=1.0, foreloebig=0.0)]
        dok = (b or {}).get("dokumenter") or [None]
        for x in linjer:
            ark.raekke([
                f["fund_code"], f["fund_name"], (b or {}).get("gennem"),
                GRUNDLAG.get(b["grundlag"], b["grundlag"]) if b else "Ingen fordeling",
                (b or {}).get("dato"), f["nav_dato"], x["top"], x["under"], x["under_kode"],
                x["andel"], f["n"].get("nav_dkk"),
                lambda rad, k: f"={k['Andel']}{rad}*{k['Fondens NAV DKK']}{rad}",
                x["foreloebig"], kilder.id_(dok[0]), kilder.fil(dok[0])])
    forste, sidste = 2, ark.rad
    ark.raekke([None, "I alt", None, None, None, None, None, None, None, None, None,
                lambda rad, k: f"=SUM({k['Beløb DKK']}{forste}:{k['Beløb DKK']}{sidste})",
                None, None, None], fed=True)
    ark.afslut()


def _vaerdibevaegelser(wb, r):
    vb = r.get("vaerdibevaegelser")
    if not vb:
        return
    ark = Ark(wb, "Værdibevægelser", [
        ("Navn", 34, None, None), ("Via", 34, None, None), ("Kvartalet DKK", 14, "tal", None),
        ("År til dato DKK", 14, "tal", None), ("Bemærkning", 70, None, None)])
    for x in sorted(vb["linjer"], key=lambda x: -(x.get("kvartal") or 0)):
        ark.raekke([x["navn"], x["fond"], x.get("kvartal"), x.get("aar"),
                    "Værdiskabt, andel (fanen Beholdninger): ultimo minus primo, gange DKK-kursen "
                    "på rapportdatoen"])
    for p, liste in vb["usammenlignelige"].items():
        for fond, navn in liste:
            ark.raekke([navn, fond, None, None, f"Kan ikke sammenlignes for {PERIODER[p].lower()}: "
                        "afhændet, omdøbt eller købt før perioden"])
    for p, fonde in vb["uaendrede"].items():
        for fond in fonde:
            ark.raekke([None, fond, None, None, f"Ingen nye beholdningstal for {PERIODER[p].lower()}"])
    ark.afslut()


def _valutakurser(wb, r):
    ark = Ark(wb, "Valutakurser", [
        ("Valuta", 8, None, None), ("Ønsket dato", 12, "dato", None),
        ("DKK pr. enhed", 12, "kurs", None), ("Kursdato", 12, "dato", None),
        ("Kilde", 60, None, "kilde")])
    for (valuta, d), k in sorted(r["kursopslag"].items(), key=lambda x: (x[0][0], x[0][1])):
        kurs, kursdato = k or (None, None)
        ark.raekke([valuta, d, kurs, kursdato,
                    "DKK, kurs 1" if valuta == "DKK" else
                    "Danmarks Nationalbank (dbo.fx_rates_nationalbanken), seneste kurs på eller før"])
    ark.afslut()


def _stamdata(wb, r):
    ark = Ark(wb, "Stamdata", [
        ("Fondskode", 11, None, None), ("Fond", 34, None, None),
        ("Juridisk navn", 38, None, None), ("Forvalter", 28, None, None),
        ("Geografi", 20, None, None), ("Investeringsstrategi", 22, None, None), ("Vintage", 8, None, None),
        ("Løbetid", 9, None, None), ("Valuta", 7, None, None), ("Benchmark", 20, None, None),
        ("Kilde", 60, None, "kilde")])
    for f in r["fonde"]:
        ark.raekke([f["fund_code"], f["fund_name"], f.get("juridisk_navn", f["fund_name"]),
                    f["manager"], f["geografi"], f["segment"],
                    f["vintage"], f["loebetid"], f["valuta"], f["benchmark"],
                    "rpt.v_fund_master (geografi, investeringsstrategi, løbetid, forvalter); "
                    "config/fund_info.yaml (beskrivelse, benchmark)"])
    ark.afslut()


def _laes_mig(wb, r, faner):
    ws = wb.create_sheet("Læs mig", 0)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    linjer = [
        ("Dataudtræk", f"{r['navn']} – fondsrapport pr. {r['as_of']:%d.%m.%Y}"),
        ("Dannet", f"{datetime.now():%d.%m.%Y %H:%M}"),
        ("Fortroligt", "Til intern afstemning; indeholder dokumentstier og interne id'er. "
                       "Sendes ikke til klienten."),
        ("Kildedokumenter", str(DOKUMENTER)),
        ("", ""),
        ("Sådan afstemmes", "Hver række fra databasen har et Dokument-id (link til fanen "
                            "Dokumenter) og et link til PDF'en. Brune kolonner er kilder, "
                            "terrakotta kolonner er kontroller, der regner tallet igen fra "
                            "rækkerne i de andre faner. Afledte tal (resttilsagn, PICC, DPI, "
                            "TVPI, ejerandel, værdiændring, DKK-beløb) er Excel-formler."),
        ("", ""),
    ] + [(navn, tekst) for navn, tekst in faner] + [("", "")]
    if not r["datotjek"]["ens"]:
        for x in r["datotjek"]["foraeldede"]:
            linjer.append(("Advarsel", f"{x['fond']} er opgjort pr. {x['dato']:%d.%m.%Y}, "
                                       "før rapportdatoen."))
        for x in r["datotjek"]["skaeve_beholdninger"]:
            linjer.append(("Advarsel", f"{x['fond']}: beholdninger pr. "
                                       f"{x['beholdninger_dato']:%d.%m.%Y}, nøgletal pr. "
                                       f"{x['nav_dato']:%d.%m.%Y}."))
    for i, (a, b) in enumerate(linjer, 1):
        ws.cell(i, 1, a).font = Font(bold=True, color=style.PPIM_SAGE_DARK.lstrip("#"))
        c = ws.cell(i, 2, b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(1, 2).font = Font(bold=True, size=14)


FANER = [
    ("Nøgletal", "Én række pr. fond: tilsagn, indbetalt, udloddet, NAV og de afledte nøgletal, "
                 "med kontrol mod kapitalstrømmene og IRR regnet igen med XIRR."),
    ("Udvikling", "Kvartalet og år til dato pr. fond og for porteføljen: fra NAV primo til "
                  "ultimo, i fondens valuta og i DKK."),
    ("Værdibro", "Fra indbetalt kapital til NAV pr. fond, med kontrolrække."),
    ("NAV", "Alle NAV-tal brugt i rapporten, pr. fond og kvartal."),
    ("Kapitalstrømme", "Alle kapitalkald, udlodninger, udligninger og kildeskat, inkl. skønnede "
                       "kapitalkald; kolonnen I nøgletal viser, hvad der indgår."),
    ("Kapitalkonto", "Forvalternes kapitalkontoopgørelser (gebyrer, carry, forvalterens tal)."),
    ("Tilsagn", "Tilsagnshændelser pr. fond."),
    ("IRR", "IRR-cash flows pr. fond, med NAV som sidste cash flow; grundlaget for XIRR-"
            "kontrollen."),
    ("Beholdninger", "Fondenes beholdninger ultimo og primo, hele fonden og investors andel. "
                     "For feederne (Gennem udfyldt) er det den underliggende fonds positioner, og "
                     "ejerandelen er feederens NAV / den underliggende fonds samlede dagsværdi."),
    ("Værdibevægelser", "Beholdningernes værdibevægelse i kvartalet og året."),
    ("Brancher", "Siden Branchefordeling: hver fonds GICS-andele (fra rpt.v_industry_exposure) "
                 "ganget med fondens NAV i DKK; Foreløbig er den del af linjen, hvis klassifikation "
                 "ikke er gennemgået (is_suggested)."),
    ("Geografi", "Siden Geografisk eksponering: hver fonds andele pr. region og delregion (fra "
                 "rpt.v_geography_exposure) ganget med fondens NAV i DKK; Foreløbig som i "
                 "Brancher."),
    ("Valutakurser", "Alle DKK-kurser, rapporten har brugt."),
    ("Stamdata", "Geografi, investeringsstrategi, løbetid og forvalter pr. fond."),
    ("Dokumenter", "Alle kildedokumenter med type, periode, godkendelse og link til PDF'en; "
                   "Pladsholder = Ja er en manuelt oprettet post uden originalt dokument."),
]


def skriv(r):
    """Projektmappen med rapportens data og kilder. Returnerer en openpyxl Workbook."""
    wb = Workbook()
    wb.remove(wb.active)
    kilder = Kilder(r["dokumenter"])
    # Dokumenter først, så de andre fanes dokument-id'er kan linke til rækken
    _dokumenter(wb, r, kilder)
    ks = _kapitalstroemme(wb, r, kilder)
    nav_rad = _nav(wb, r, kilder)
    _kapitalkonto(wb, r, kilder)
    _tilsagn(wb, r, kilder)
    irr_blokke = _irr(wb, r)
    _noegletal(wb, r, kilder, ks, nav_rad, irr_blokke)
    _vaerdibro(wb, r)
    _udvikling(wb, r, ks, nav_rad)
    _beholdninger(wb, r, kilder)
    _vaerdibevaegelser(wb, r)
    _eksponering(wb, r, kilder, "brancher")
    _eksponering(wb, r, kilder, "regioner")
    _valutakurser(wb, r)
    _stamdata(wb, r)
    _laes_mig(wb, r, [(n, t) for n, t in FANER if n in wb.sheetnames])
    raekkefoelge = ["Læs mig"] + [n for n, _ in FANER if n in wb.sheetnames]
    wb._sheets = [wb[n] for n in raekkefoelge]
    wb.active = 0
    return wb
