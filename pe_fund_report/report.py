"""Samler investorrapporten: forside, porteføljeoversigt, én side pr. fond og slut-slide."""
import random
from contextlib import nullcontext
from pathlib import Path

from pptx import Presentation

from . import anonym as anonymisering
from . import data, excel, log, slides
from .deck import L_FORSIDE, L_SKILLE, L_SLUT, _find_master, _ny_slide, _ryd_slides, _saet_ph
from .paths import EXPORTS, MASTER


def byg(investor_code, as_of=None, outfile=None, master=None, med_data=True, anonym=False,
        med_supplerende=True):
    """Byg rapporten for én investor og returnér stien til pptx-filen.

    as_of: rapportdato (ÅÅÅÅ-MM-DD); standard er investors seneste NAV-dato.
    outfile: standard er exports/<INVESTOR>_<dato>.pptx (ANONYM_<dato>.pptx med anonym).
    med_data: gem også dataudtrækket <outfile>_data.xlsx med kilder til hvert tal.
    med_supplerende: gem også <outfile>_supplerende_slides.pptx med de slides, der er taget ud
    af rapporten (_byg_supplerende).
    anonym: uden navne, logoer, beskrivelser, fondsstørrelse og ejerandel, og alle beløb
    ganget med én skjult tilfældig faktor; aldrig med dataudtræk. Stopper uden at gemme, hvis
    et oprindeligt navn står i decket.

    Hver bygning skrives i log/ (log.byggelog): hvem, hvornår, hvilken kode, filerne og alt,
    hvad den skriver i terminalen - også når den fejler.
    """
    with log.byggelog(investor_code, as_of, anonym) as lg:
        r = rigtig = lg.r = data.investor_rapport(investor_code, as_of)
        skala = 1.0
        if anonym:
            skala = _skjult_faktor()
            r, originaler = anonymisering.anonymiser(rigtig, skala)
            med_data = False
        with slides.beloebsskala(skala) if anonym else nullcontext():
            prs = _byg_deck(r, master)
            ekstra = _byg_supplerende(r, master) if med_supplerende else None
        if anonym:
            # begge decks tjekkes, før noget gemmes
            laek = sorted({x for p in (prs, ekstra) if p for x in anonymisering.laek(p, originaler)})
            if laek:
                raise ValueError("Den anonymiserede udgave nævner stadig: " + ", ".join(laek)
                                 + " - intet er gemt")

        navn = "ANONYM" if anonym else r["investor"]
        oenske = Path(outfile or EXPORTS / f"{navn}_{r['as_of']:%Y-%m-%d}.pptx")
        outfile = _gem(prs, oenske)
        lg.filer.append(outfile)
        print(f"gemt: {outfile}  ({len(r['fonde'])} fonde, {len(prs.slides._sldIdLst)} slides"
              + (", anonymiseret)" if anonym else ")"))
        if ekstra:
            # navngivet efter det ønskede navn, ikke et _v2, fordi rapporten var åben
            sti = _gem(ekstra, oenske.with_name(f"{oenske.stem}_supplerende_slides.pptx"))
            lg.filer.append(sti)
            print(f"gemt: {sti}  (supplerende slides, {len(ekstra.slides._sldIdLst)} slides)")
        if med_data:
            # dataudtrækket ved siden af decket: samme navn med _data, til intern afstemning
            xlsx = _gem(excel.skriv(r), outfile.with_name(f"{outfile.stem}_data.xlsx"))
            lg.filer.append(xlsx)
            print(f"gemt: {xlsx}  (data og kildedokumenter)")
        # advarslerne er til bygningen, ikke decket, så de nævner de rigtige fonde
        _advarsler(rigtig)
    return outfile


def _skjult_faktor():
    """En tilfældig faktor til beløbene, holdt væk fra 1, så de rigtige tal ikke skinner igennem."""
    rng = random.SystemRandom()
    return rng.uniform(0.55, 0.8) if rng.random() < 0.5 else rng.uniform(1.25, 1.8)


def _byg_deck(r, master=None):
    """Hele præsentationen i slide-rækkefølgen; gemmes ikke."""
    prs = Presentation(str(_find_master(master or MASTER)))
    _ryd_slides(prs)

    s = _ny_slide(prs, L_FORSIDE)
    # porteføljen rummer også private credit og ejendomme, så ikke "private equity"
    _saet_ph(s, 0, f"{r['navn']} – alternative investeringer")
    _saet_ph(s, 1, f"Porteføljerapport pr. {r['as_of']:%d.%m.%Y}")
    slides.sammendrag(prs, r)
    slides.udvikling(prs, r)            # springes over uden DKK-kurser for alle fonde

    s = _ny_slide(prs, L_SKILLE[0])
    _saet_ph(s, 0, "Porteføljeoversigt")
    slides.tilsagn_paa_tvaers(prs, r)
    slides.aendringer_pr_fond(prs, r)
    slides.valutafordeling(prs, r)      # springes over uden DKK-kurser for alle fonde
    # geografien efter tilsagn, cirkeldiagrammerne og ejendomme for sig står i de supplerende
    # slides (_byg_supplerende); Geografisk eksponering og treemaps afløser dem her
    slides.fordeling(prs, r, "segment", "Porteføljefordeling efter investeringsstrategi")
    slides.branchekort(prs, r)          # hele porteføljen, ejendommene som Real Estate
    slides.branchekort_aktivklasser(prs, r)     # Private Equity, Private Credit og ejendomme
    slides.regionskort(prs, r)          # hele porteføljen
    slides.regionskort_aktivklasser(prs, r)     # Private Equity, Private Credit og ejendomme
    slides.stoerste_positioner(prs, r)  # selskaber, underliggende fonde og ikke-gennemlyste fonde
    # i fondens valuta, så den springes over, hvis fondene har forskellige valutaer; Største
    # positioner ovenfor er den samme gennemlysning i DKK
    slides.gennemlysning(prs, r)
    # Største værdibevægelser står i de supplerende slides og i dataudtrækkets fane
    # Værdibevægelser

    s = _ny_slide(prs, L_SKILLE[1])
    _saet_ph(s, 0, "Fondsoversigter")
    slides.portefoeljeoversigt(prs, r)  # oversigten over fondene indleder fondssiderne
    for f in r["fonde"]:
        slides.fondsside(prs, f, r["navn"])
        slides.virksomhedsprofiler(prs, f, r["navn"])   # kun fonde med porteføljeselskaber

    s = _ny_slide(prs, L_SKILLE[2])
    _saet_ph(s, 0, "Appendiks")
    slides.definitioner(prs, r)
    slides.gics_oversigt(prs, r)        # GICS-sektorerne og deres branchegrupper
    _ny_slide(prs, L_SLUT)
    slides.sidefod(prs, r["navn"])      # til sidst, så sidetallene er endelige
    return prs


def _byg_supplerende(r, master=None):
    """De supplerende slides: dem, der er taget ud af rapporten, men stadig bygges af de samme
    data. En slide, der fjernes fra rapporten, flyttes hertil. Hver springer sig selv over uden
    sine data (fx uden DKK-kurser eller ejendomme); gemmes ikke."""
    prs = Presentation(str(_find_master(master or MASTER)))
    _ryd_slides(prs)

    s = _ny_slide(prs, L_FORSIDE)
    _saet_ph(s, 0, f"{r['navn']} – supplerende slides")
    _saet_ph(s, 1, f"Slides uden for porteføljerapporten pr. {r['as_of']:%d.%m.%Y}")
    slides.lange_tilsagn(prs, r)
    slides.fordeling(prs, r, "geografi", "Geografisk fordeling efter tilsagn")
    slides.brancher(prs, r)             # cirkeldiagrammerne
    slides.regioner(prs, r)
    slides.branchekort_ejendomme(prs, r)        # ejendommene for sig
    slides.regionskort_ejendomme(prs, r)
    slides.vaerdibevaegelser(prs, r)

    s = _ny_slide(prs, L_SKILLE[2])
    _saet_ph(s, 0, "Appendiks")
    slides.definitioner(prs, r)
    _ny_slide(prs, L_SLUT)
    slides.sidefod(prs, r["navn"])
    return prs


def _advarsler(r):
    """Forhold, der skal ses på før udsendelse; skrives ved bygningen, ikke i decket."""
    _advar_om_datoer(r["datotjek"])
    for f in r["fonde"]:
        for advarsel in f.get("advarsler", []):
            print(f"ADVARSEL: {f['fund_name']}: {advarsel}")
    _advar_om_klassifikation(r["brancher"], "branchen")
    _advar_om_klassifikation(r["regioner"], "geografien")
    # forskellige valutaer regnes sammen i DKK; kun en manglende kurs efterlader hul i totalerne
    uden_dkk = [f["fund_name"] for f in r["fonde"] if f["n"].get("nav_dkk") is None]
    if r["valuta"] is None and uden_dkk:
        print(f"ADVARSEL: ingen DKK-kurs for {', '.join(uden_dkk)} - fondene har forskellige "
              "valutaer, så porteføljens tal i DKK, fordelingerne og IRR i DKK mangler")


MAKS_VERSIONER = 50


def _gem(prs, outfile):
    """Gem decket (eller projektmappen; begge har .save) og returnér den sti, det endte på.

    PowerPoint låser filen, mens den er åben. I stedet for at fejle lægges et versionsnummer
    på - MOTORTRAMP_2026-06-30_v2.pptx og så videre - så en bygning aldrig går tabt, fordi
    nogen sad med den forrige udgave åben.
    """
    outfile.parent.mkdir(parents=True, exist_ok=True)
    for version in range(1, MAKS_VERSIONER + 1):
        sti = (outfile if version == 1
               else outfile.with_name(f"{outfile.stem}_v{version}{outfile.suffix}"))
        try:
            prs.save(str(sti))
        except PermissionError:
            print(f"{sti.name} er låst (åben i PowerPoint?) - prøver næste versionsnummer")
            continue
        return sti
    raise PermissionError(f"Kunne ikke skrive {outfile.name} i {MAKS_VERSIONER} versioner - "
                          "luk de åbne udgaver i PowerPoint")


def _advar_om_klassifikation(b, hvad):
    """Skriv den del af NAV, hvis branche eller geografi ikke er gennemgået, så den kan rettes
    før udsendelse."""
    if b and b["foreloebig_andel"] and b["foreloebig_andel"] >= 0.0005:
        print(f"ADVARSEL: {hvad} for {b['foreloebig_andel']:.1%} af NAV er en foreløbig "
              f"klassifikation (is_suggested) - {', '.join(b['foreloebige_fonde'])}; "
              "se fanerne Brancher og Geografi i dataudtrækket")


def _advar_om_datoer(tjek):
    """Skriv uens opgørelsesdatoer ud ved bygningen, så de ses før decket sendes videre."""
    if tjek["ens"]:
        print(f"datotjek: alle fonde er opgjort pr. {tjek['rapport_dato']:%d.%m.%Y}")
        return
    for x in tjek["foraeldede"]:
        print(f"ADVARSEL: {x['fond']} er opgjort pr. {x['dato']:%d.%m.%Y}, "
              f"{x['kvartaler']} kvartal(er) før rapportdatoen "
              f"{tjek['rapport_dato']:%d.%m.%Y} - totaler blander opgørelsesdatoer")
    for x in tjek["skaeve_beholdninger"]:
        print(f"ADVARSEL: {x['fond']} har beholdninger pr. "
              f"{x['beholdninger_dato']:%d.%m.%Y}, men nøgletal pr. {x['nav_dato']:%d.%m.%Y}")
