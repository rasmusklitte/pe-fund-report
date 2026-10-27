"""Aletas tal for investorens fonde: udtrækkene i data/aleta/<INVESTOR>/.

Aleta er investorens eget system med de samme fonde. noegletal_<dato>.xlsx har én række pr. fond
(Aletas beløb i DKK, undtagen nav i fondens valuta), dateret den første i måneden efter
rapportdatoen (aleta_dato()); private_markets_transactions_*.xlsx (tidligere
pm_transactions_*.xlsx) er Aletas transaktioner (det nyeste udtræk efter datoen i navnet). afstem_transaktioner() parrer transaktionerne med rapportens kapitalstrømme, og
irr_sammenligning() stiller rapportens IRR op mod Aletas til sliden IRR sammenlignet med Aleta.
Afstemningen kontrol for kontrol er tests/aleta.py.
"""
import io
import math
import shutil
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .data import xirr
from .paths import ALETA

BELOEB_TOLERANCE = 1.0          # to kapitalstrømme er ens, når beløbene er højst 1 fra hinanden
MAX_DAGE = 183                  # samme beløb med højst så mange dages forskel er samme strøm


def aleta_dato(as_of):
    """Aletas dato for en rapport pr. as_of: den første i den følgende måned (reglen)."""
    return date(as_of.year + as_of.month // 12, as_of.month % 12 + 1, 1)


def find(investor, as_of, rod=ALETA):
    """(nøgletallene, transaktionerne) for investoren pr. as_of, eller None uden et udtræk på
    Aletas dato. Transaktionerne er det nyeste udtræk (navnet bærer udtræksdatoen), None uden."""
    mappe = rod / investor
    noegletal = mappe / f"noegletal_{aleta_dato(as_of):%Y-%m-%d}.xlsx"
    if not noegletal.exists():
        return None
    filer = [*mappe.glob("pm_transactions_*.xlsx"),
             *mappe.glob("private_markets_transactions_*.xlsx")]
    return noegletal, max(filer, key=_udtraeksdato, default=None)


def _udtraeksdato(sti):
    """Sorteringsnøglen for et transaktionsudtræk: datoen sidst i navnet (»…_2026-10-26«), så
    navnets præfiks ikke afgør, hvilket der er nyest; uden dato filens ændringstid."""
    try:
        dato = date.fromisoformat(sti.stem[-10:])
    except ValueError:
        dato = date.min
    return dato, sti.stat().st_mtime


# Aletas securityCode -> vores fund_code; navnene ligner ikke altid hinanden (Park Square er vores
# feeder PPPC2, Newbury PPPE2, "API Property DK" er CRED)
FONDE = {
    "APIPropertyDK(CRED)": "CRED",
    "KAPE2020": "KAPE20",
    "KAPE2021": "KAPE21",
    "KAPE2022": "KAPE22",
    "Seed Capital": "SEED4",
    "Home.Earth": "HOME",
    "Private Debt III": "PD3LGT",
    "Catacap III": "CATACAP3",
    "Park Square - Private Credit II": "PPPC2",
    "Newbury - Private Equity II": "PPPE2",
}
# Aletas fondsnavne i transaktionsudtrækket -> vores fund_code
TRANSAKTIONSFONDE = {
    "API Property DK (CRED)": "CRED",
    "CataCap III K/S": "CATACAP3",
    "Home.Earth K/S": "HOME",
    "K Alternativ Private Equity 2020 K/S": "KAPE20",
    "K Alternativ Private Equity 2021 K/S": "KAPE21",
    "K Alternativ Private Equity 2022 K/S": "KAPE22",
    "Newbury - Private Equity II": "PPPE2",
    "Park Square - Private Credit II": "PPPC2",
    "Private Debt III LGT K/S": "PD3LGT",
    "SEED Capital Denmark IV K/S": "SEED4",
}


def _projektmappe(sti):
    """Udtrækket som projektmappe. Excel låser filen, mens den er åben, så Python ikke kan åbne
    den; Windows kan stadig kopiere den, så låst læses en midlertidig kopi i stedet for at fejle."""
    try:
        return load_workbook(sti, read_only=True, data_only=True)
    except PermissionError:
        with tempfile.TemporaryDirectory() as mappe:
            indhold = Path(shutil.copy2(sti, mappe)).read_bytes()
        return load_workbook(io.BytesIO(indhold), read_only=True, data_only=True)


def laes_aleta(sti):
    """Aletas rækker som dicts med kolonnenavnene som nøgler; vintage som årstal."""
    ws = _projektmappe(sti).worksheets[0]
    raekker = ws.iter_rows(values_only=True)
    hoved = next(raekker)
    ud = []
    for r in raekker:
        if not any(r):
            continue
        x = dict(zip(hoved, r, strict=True))
        if hasattr(x["vintage"], "year"):
            x["vintage"] = x["vintage"].year
        ud.append(x)
    return ud


def laes_transaktioner(sti):
    """Aletas transaktioner (PM Transactions) som dicts: fond (vores fund_code), id, type ("kald"
    eller "udlodning"), dato (value date), beloeb (positivt, i fondens valuta), beloeb_dkk, kurs,
    i_noegletal (Aletas »Affects Multiples«), vintage (fondens startdato i Aleta), kommentar."""
    ws = _projektmappe(sti)["PM Transactions"]
    raekker = ws.iter_rows(values_only=True)
    hoved = next(raekker)
    ud = []
    for r in raekker:
        if not any(r):
            continue
        x = dict(zip(hoved, r, strict=True))
        ud.append(dict(fond=TRANSAKTIONSFONDE.get(x["Fund"], x["Fund"]), id=x["Transaction ID"],
                       type="kald" if x["Direction"] == "Drawdown" else "udlodning",
                       dato=date.fromisoformat(str(x["Value Date"])[:10]),
                       beloeb=float(x["Amount (fund CCY)"]), beloeb_dkk=float(x["Amount DKK"]),
                       kurs=float(x["FX fund→DKK (Aleta)"]), i_noegletal=bool(x["Affects Multiples"]),
                       vintage=date.fromisoformat(str(x["Vintage"])[:10]),
                       kommentar=x["Comment"] or ""))
    return ud


# --- transaktionerne ------------------------------------------------------------------------------

def _rapportens_stroemme(f):
    """Rapportens kapitalstrømme samlet pr. (dato, type): beløb (positivt), kurs og reference."""
    cf = f.get("kilder", {}).get("cash_flows")
    kurser, refs, kaldsbreve = {}, {}, {}
    if cf is not None:
        for d, k, ref, nd in zip(cf.value_date, cf.fx_rate, cf.flow_ref, cf.notice_date,
                                 strict=True):
            if k is not None and not math.isnan(k):
                kurser[d] = float(k)
            if ref:
                refs.setdefault(d, []).append(ref)
            if isinstance(nd, date):
                kaldsbreve[d] = date.fromisoformat(str(nd)[:10])
    if f["valuta"] == "DKK":
        kurser = {}
    samlet = {}
    for d, b, t in f["kapitalstroemme"]:
        samlet[d, t] = samlet.get((d, t), 0.0) + (-b if t == "kald" else b)
    return [{"dato": d, "type": t, "beloeb": b,
             "kurs": 1.0 if f["valuta"] == "DKK" else kurser.get(d),
             "ref": ", ".join(dict.fromkeys(refs.get(d, []))), "kaldsbrev": kaldsbreve.get(d)}
            for (d, t), b in sorted(samlet.items()) if abs(b) >= BELOEB_TOLERANCE]


def _par(aleta, rapport):
    """Par Aletas og rapportens strømme (samme fond og type): [(art, a, r)] med a eller r None for
    en strøm uden makker. Først samme dato og beløb, så samme beløb inden for MAX_DAGE (nærmeste
    dato), så samme dato og til sidst nærmeste dato inden for en måned med et andet beløb."""
    a_rest, r_rest, par = list(aleta), list(rapport), []

    def tag(art, betingelse, afstand):
        for a in sorted(a_rest, key=lambda a: a["dato"]):
            kandidater = [r for r in r_rest if betingelse(a, r)]
            if kandidater:
                r = min(kandidater, key=lambda r: afstand(a, r))
                par.append((art, a, r))
                a_rest.remove(a)
                r_rest.remove(r)

    dage = lambda a, r: abs((r["dato"] - a["dato"]).days)
    samme_beloeb = lambda a, r: abs(a["beloeb"] - r["beloeb"]) <= BELOEB_TOLERANCE
    tag("ens", lambda a, r: a["dato"] == r["dato"] and samme_beloeb(a, r), dage)
    tag("anden dato", lambda a, r: samme_beloeb(a, r) and dage(a, r) <= MAX_DAGE, dage)
    tag("andet beløb", lambda a, r: a["dato"] == r["dato"], dage)
    tag("andet beløb", lambda a, r: dage(a, r) <= 31,
        lambda a, r: (dage(a, r), abs(a["beloeb"] - r["beloeb"])))
    par += [("kun i Aleta", a, None) for a in a_rest] + [("kun i rapporten", None, r) for r in r_rest]
    return sorted(par, key=lambda p: (p[1] or p[2])["dato"])


# status pr. art, som i kontrollerne; strømme efter rapportdatoen er kun til orientering
TRANSAKTIONSSTATUS = {"ens": "ok", "anden dato": "lille afvigelse", "andet beløb": "afviger",
                      "kun i Aleta": "afviger", "kun i rapporten": "afviger",
                      "efter rapportdatoen": "efter rapportdatoen"}


def _kr(x):
    return f"{x:,.0f}".replace(",", ".")


def afstem_transaktioner(transaktioner, r, as_of):
    """Én linje pr. kapitalstrøm: Aletas transaktioner (laes_transaktioner()) parret med
    rapportens kapitalstrømme pr. fond, type og dato. art er ens, anden dato, andet beløb, kun i
    Aleta, kun i rapporten eller efter rapportdatoen (Aletas strømme efter as_of, som rapporten
    ikke læser); status er kontrollernes ok / lille afvigelse / afviger. Aletas transaktioner på
    samme dato og af samme type lægges sammen, ligesom rapportens konti. kaldsbrev er sand, når
    Aletas dato er kaldsbrevets; uden_for_noegletal er beløbet, Aleta holder ude af nøgletallene."""
    fonde = {f["fund_code"]: f for f in r["fonde"]}
    ud = []
    for kode in dict.fromkeys(t["fond"] for t in transaktioner):
        f = fonde.get(kode)
        egne = [t for t in transaktioner if t["fond"] == kode]
        for type_ in ("kald", "udlodning"):
            samlet = {}
            for t in (t for t in egne if t["type"] == type_):
                s = samlet.setdefault(t["dato"], {"dato": t["dato"], "type": type_,
                                                   "beloeb": 0.0, "beloeb_dkk": 0.0, "id": [],
                                                   "kommentar": [], "i_noegletal": 0.0})
                s["beloeb"] += t["beloeb"]
                s["beloeb_dkk"] += t["beloeb_dkk"]
                s["id"].append(t["id"])
                s["kommentar"] += [t["kommentar"]] if t["kommentar"] else []
                s["i_noegletal"] += t["beloeb"] if t["i_noegletal"] else 0.0
            aleta = list(samlet.values())
            senere = [a for a in aleta if a["dato"] > as_of]
            rapport = [] if f is None else [x for x in _rapportens_stroemme(f) if x["type"] == type_]
            par = _par([a for a in aleta if a["dato"] <= as_of], rapport)
            par += [("efter rapportdatoen", a, None) for a in senere]
            for art, a, rr in par:
                linje = {"fond": kode, "type": type_, "art": art,
                         "status": TRANSAKTIONSSTATUS[art],
                         "valuta": None if f is None else f["valuta"],
                         "aleta_dato": a and a["dato"], "rapport_dato": rr and rr["dato"],
                         "dage": None if a is None or rr is None else (rr["dato"] - a["dato"]).days,
                         "aleta": a and a["beloeb"], "rapport": rr and rr["beloeb"],
                         "aleta_dkk": a and a["beloeb_dkk"],
                         "aleta_kurs": a and a["beloeb_dkk"] / a["beloeb"],
                         "rapport_kurs": rr and rr["kurs"], "ref": rr["ref"] if rr else "",
                         "aleta_id": ", ".join(a["id"]) if a else "",
                         "vintage": egne[0]["vintage"], "bemaerkning": "",
                         "kaldsbrev": bool(art == "anden dato" and rr["kaldsbrev"] == a["dato"]),
                         "uden_for_noegletal": a["beloeb"] - a["i_noegletal"] if a else 0.0}
                linje["rapport_dkk"] = (None if rr is None or rr["kurs"] is None
                                        else rr["beloeb"] * rr["kurs"])
                linje["afvigelse"] = (None if a is None or rr is None
                                      else rr["beloeb"] - a["beloeb"])
                linje["afvigelse_dkk"] = (None if a is None or linje["rapport_dkk"] is None
                                          else linje["rapport_dkk"] - a["beloeb_dkk"])
                noter = list(a["kommentar"]) if a else []
                if linje["kaldsbrev"]:
                    noter.append("Aletas dato er meddelelsesdatoen (notice date); rapporten "
                                 "bruger betalingsdatoen")
                if linje["uden_for_noegletal"] > BELOEB_TOLERANCE:
                    noter.append(f"{_kr(linje['uden_for_noegletal'])} tæller ikke i Aletas "
                                 "nøgletal (equalization)")
                if f is None:
                    noter.append("fonden er ikke i rapporten")
                linje["bemaerkning"] = "; ".join(noter)
                ud.append(linje)
    return ud


def _aletas_irr(tx):
    """Aletas IRR genberegnet: XIRR over Aletas transaktioner til og med rapportdatoen i DKK."""
    return [(x["aleta_dato"], (-1 if x["type"] == "kald" else 1) * x["aleta_dkk"])
            for x in tx if x["aleta_dato"] is not None and x["art"] != "efter rapportdatoen"]


def _ens(a, b, graense=1.0):
    return a is not None and b is not None and abs(a - b) < graense


def _gp_konto(f, kode):
    """Forvalterens kapitalkonto (ITD) for kode pr. fondens NAV-dato, ellers None."""
    ka = f.get("kilder", {}).get("kapitalkonto")
    if ka is None or ka.empty:
        return None
    x = ka[(ka.account_code == kode) & (ka.basis == "ITD") & (ka.period_end == f["nav_dato"])]
    return float(x.amount.iloc[0]) if not x.empty else None


def _uden_nav(f):
    """Rapportens IRR-strømme uden NAV'en som sidste strøm."""
    return [(d, b) for d, b in f["irr_flows"] if (d, b) != (f["nav_dato"], f["n"]["nav"])]


def _irr_dkk(f, nav_dkk, nav_dato, datoer=None):
    """Rapportens IRR i DKK: hver strøm med kursen på sin dato (kapitalstrømmenes fx_rate),
    nav_dkk pr. nav_dato. datoer {rapportens dato: en anden dato} flytter strømmene (kursen
    følger med fra rapportens dato). None, hvis en kurs mangler."""
    cf = f.get("kilder", {}).get("cash_flows")
    if cf is None:
        return None
    kurser = {d: float(k) for d, k in zip(cf.value_date, cf.fx_rate, strict=True)
              if k is not None and not math.isnan(k)}
    stroemme = _uden_nav(f)
    if any(d not in kurser for d, _ in stroemme):
        return None
    datoer = datoer or {}
    return xirr([(datoer.get(d, d), b * kurser[d]) for d, b in stroemme] + [(nav_dato, nav_dkk)])


# --- IRR mod Aleta --------------------------------------------------------------------------------

def _stroemme(antal):
    """Antallet til årsagerne på sliden, hvor kapitalstrømmene hedder betalinger."""
    return f"{antal} betaling" if antal == 1 else f"{antal} betalinger"


def irr_sammenligning(f, a, tx, dato, kurs):
    """Fondens IRR mod Aletas: {irr, irr_dkk, irr_aletas_datoer, aleta, forskel, aarsager,
    metode_ok, dato}. a er Aletas række, tx fondens linjer fra afstem_transaktioner(), dato
    Aletas dato og kurs fondens DKK-kurs pr. dato. irr_dkk er rapportens strømme i DKK med kursen
    på hver strøms dato og NAV omregnet pr. dato (i DKK-fonde rapportens IRR); irr_aletas_datoer
    de samme strømme flyttet til Aletas datoer. forskel er rapportens IRR i DKK (irr_dkk) -
    Aletas, så begge er i DKK. aarsager er korte
    sætninger uden beløb, hver kun når tallene viser den; metode_ok, at Aletas IRR er XIRR over
    Aletas egne transaktioner med markedsværdien pr. dato som sidste strøm."""
    n, dkk = f["n"], f["valuta"] == "DKK"
    tx = [t for t in tx if t["art"] != "efter rapportdatoen"]
    flyt = {t["rapport_dato"]: t["aleta_dato"] for t in tx if t["art"] == "anden dato"}
    if dkk:
        irr_dkk = n["irr"]
        flyttet = xirr([(flyt.get(d, d), b) for d, b in _uden_nav(f)] + [(dato, n["nav"])])
    else:
        nav_dkk = None if kurs is None else n["nav"] * kurs
        irr_dkk = None if nav_dkk is None else _irr_dkk(f, nav_dkk, dato)
        flyttet = None if nav_dkk is None else _irr_dkk(f, nav_dkk, dato, flyt)
    aleta_irr = a.get("irr")
    metode_ok = bool(tx) and _ens(xirr(_aletas_irr(tx) + [(dato, a["marketValue"])]),
                                  aleta_irr, 0.00005)

    aarsager = []
    tidligere = [t for t in tx if t["art"] == "anden dato" and t["dage"] > 0]
    if tidligere:
        aarsager.append(f"{_stroemme(len(tidligere))} har en tidligere dato hos Aleta")
    senere = [t for t in tx if t["art"] == "anden dato" and t["dage"] < 0]
    if senere:
        aarsager.append(f"{_stroemme(len(senere))} har en senere dato hos Aleta")
    for t in (t for t in tx if t["status"] == "afviger"):
        hvad = "kaldet" if t["type"] == "kald" else "udlodningen"
        if t["art"] == "andet beløb":
            aarsager.append(f"{hvad} pr. {t['rapport_dato']:%d.%m.%Y} har et andet beløb hos Aleta")
        elif t["art"] == "kun i Aleta":
            aarsager.append(f"{hvad} pr. {t['aleta_dato']:%d.%m.%Y} findes kun hos Aleta")
        else:
            aarsager.append(f"{hvad} pr. {t['rapport_dato']:%d.%m.%Y} mangler hos Aleta")
    ca890, ca900 = _gp_konto(f, "CA890"), _gp_konto(f, "CA900")
    if ca900 is not None and _ens(a.get("nav"), ca890) and not _ens(ca890, ca900):
        aarsager.append("Aleta bruger NAV før carried interest")
    if n.get("udligning") and any(t["uden_for_noegletal"] > BELOEB_TOLERANCE for t in tx):
        aarsager.append("Aleta regner equalization ud")
    if not dkk and any(t["afvigelse_dkk"] is not None and abs(t["afvigelse_dkk"]) > 1
                       and t["status"] != "afviger" for t in tx):
        aarsager.append("Aleta omregner med sine egne valutakurser")
    return dict(irr=n["irr"], irr_dkk=irr_dkk, irr_aletas_datoer=flyttet, aleta=aleta_irr,
                forskel=None if None in (irr_dkk, aleta_irr) else irr_dkk - aleta_irr,
                aarsager=aarsager, metode_ok=metode_ok, dato=dato)


# --- nøgletallene kontrol for kontrol -------------------------------------------------------------

TOLERANCE = 0.005               # som data.GP_TOLERANCE
OK_GRAENSE = 0.01               # under 1 % af den tilladte afvigelse er ok (grøn), ikke gul


def _rvpi(n):
    return n["nav"] / n["indbetalt"] if n["indbetalt"] else None


def _irr_i_dkk(f):
    """Rapportens IRR i DKK fra f["aleta"]; uden den IRR i fondens valuta (ens for DKK-fonde)."""
    irr_dkk = (f.get("aleta") or {}).get("irr_dkk")
    return f["n"]["irr"] if irr_dkk is None else irr_dkk


# (kontrol, Aletas kolonne, rapportens tal, tolerance)
# tolerance: "eksakt", ("relativ", x), ("af_tilsagn", x) eller ("absolut", x)
KONTROLLER = [
    ("Valuta", "navCurrency", lambda f: f["valuta"], "eksakt"),
    ("Vintage", "vintage", lambda f: f["vintage"], "eksakt"),
    ("NAV", "nav", lambda f: f["n"]["nav"], ("relativ", TOLERANCE)),
    ("NAV DKK", "marketValue", lambda f: f["n"]["nav_dkk"], ("relativ", TOLERANCE)),
    ("Tilsagn DKK", "commitment", lambda f: f["n"]["tilsagn_dkk"], ("relativ", TOLERANCE)),
    ("Resttilsagn DKK", "unfundedCommitment", lambda f: f["n"]["resttilsagn_dkk"],
     ("af_tilsagn", TOLERANCE)),
    ("Indbetalt DKK", "drawdowns", lambda f: f["n"]["indbetalt_dkk"], ("relativ", TOLERANCE)),
    ("Udloddet DKK", "distributions", lambda f: f["n"]["udloddet_dkk"], ("relativ", TOLERANCE)),
    ("PICC", "calledPercent", lambda f: f["n"]["picc"], ("absolut", 0.005)),
    ("DPI", "dpi", lambda f: f["n"]["dpi"], ("absolut", 0.01)),
    ("RVPI", "rvpi", lambda f: _rvpi(f["n"]), ("absolut", 0.01)),
    ("TVPI", "tvpi", lambda f: f["n"]["tvpi"], ("absolut", 0.01)),
    # Aletas IRR er i DKK, så rapportens IRR i DKK (irr_sammenligning), som på IRR-sliden
    ("IRR", "irr", lambda f: _irr_i_dkk(f), ("absolut", 0.0025)),
]

# beløbskontrollerne -> rapportens tal i fondens valuta, til fanen Detaljer; Aletas DKK-beløb
# regnes tilbage med samme kurs (NAV er i fondens valuta hos begge)
FONDSVALUTA = {"NAV": "nav", "NAV DKK": "nav", "Tilsagn DKK": "tilsagn",
               "Resttilsagn DKK": "resttilsagn", "Indbetalt DKK": "indbetalt",
               "Udloddet DKK": "udloddet"}
# beholdningerne, som Aleta omregner med kursen på sin egen dato
BEHOLDNINGER = ("NAV DKK", "Tilsagn DKK", "Resttilsagn DKK")
STROEMME = ("Indbetalt DKK", "Udloddet DKK")


def _tal(v):
    """np.float64 og int som float, så sammenligning og Excel får almindelige tal."""
    return float(v) if isinstance(v, int | float) and not isinstance(v, bool) else v


def _afrunding(kontrol):
    """Den største afvigelse, der forsvinder i afrundingen til det viste format."""
    if kontrol in ("PICC", "IRR"):
        return 0.00005
    if kontrol in ("DPI", "RVPI", "TVPI"):
        return 0.0005
    return 0.5


def _vurder(kontrol, aleta, rapport, tolerance, tilsagn):
    """(afvigelse, tilladt afvigelse, status) for ét tal. Status er ok (ens efter afrunding, eller
    under OK_GRAENSE af den tilladte afvigelse), lille afvigelse (inden for tolerancen) eller
    afviger."""
    if aleta is None or rapport is None:
        return None, None, "afviger"
    if tolerance == "eksakt":
        return None, None, "ok" if aleta == rapport else "afviger"
    art, x = tolerance
    afvigelse = rapport - aleta
    tilladt = {"relativ": abs(aleta) * x, "af_tilsagn": abs(tilsagn or 0) * x, "absolut": x}[art]
    if abs(afvigelse) < max(_afrunding(kontrol), tilladt * OK_GRAENSE):
        return afvigelse, tilladt, "ok"
    return afvigelse, tilladt, "lille afvigelse" if abs(afvigelse) <= tilladt else "afviger"


def afstem(aleta, r, dato, kurser_aleta=None):
    """Én række pr. fond og kontrol: fond, kontrol, aleta, rapport, afvigelse, tilladt, status,
    bemaerkning og, for beløbene, valuta, kurs, kursdato og aleta_fv / rapport_fv / afvigelse_fv
    i fondens valuta. dato er Aletas dato.

    kurser_aleta er {valuta: (kurs, kursdato)} på Aletas dato. Beholdningerne (BEHOLDNINGER)
    sammenlignes som rapportens beløb i fondens valuta gange den kurs; uden den bruges rapportens
    egen DKK-omregning med rapportdatoens kurs. Indbetalt og udloddet bruger altid rapportens
    omregning med transaktionskurserne, og Kurs er da deres gennemsnit, vægtet med beløbene.
    En fond fra Aleta, der ikke er i rapporten, får status afviger i alle kontroller."""
    fonde = {f["fund_code"]: f for f in r["fonde"]}
    kurser = r.get("kurser", {})
    ud = []
    for a in aleta:
        kode = FONDE.get(a["securityCode"], a["securityCode"])
        f = fonde.get(kode)
        for kontrol, kolonne, hent, tolerance in KONTROLLER:
            linje = dict(fond=kode, aleta_navn=a["asset"], kontrol=kontrol,
                         aleta=_tal(a[kolonne]), rapport=None, afvigelse=None, tilladt=None,
                         status="afviger", bemaerkning="", valuta=None, kurs=None,
                         kursdato=None, aleta_fv=None, rapport_fv=None, afvigelse_fv=None)
            if f is None:
                linje["bemaerkning"] = "fonden er ikke i rapporten"
                ud.append(linje)
                continue
            linje["rapport"] = _tal(hent(f))
            noegle = FONDSVALUTA.get(kontrol)
            k, kdato = kurser.get(f["valuta"]) or (None, None)
            if kontrol in BEHOLDNINGER and kurser_aleta is not None:
                k, kdato = kurser_aleta.get(f["valuta"]) or (None, None)
                fv = f["n"].get(noegle)
                if k is None:
                    linje["bemaerkning"] = f"ingen {f['valuta']}-kurs pr. {dato:%d.%m.%Y}"
                linje["rapport"] = None if k is None or fv is None else _tal(fv * k)
            elif kontrol in STROEMME:
                kdato = "transaktionsdatoer"
                if f["n"].get(noegle) and f["n"].get(f"{noegle}_dkk") is not None:
                    # rapportens egen omregning med transaktionskurserne, som én gennemsnitskurs
                    k = f["n"][f"{noegle}_dkk"] / f["n"][noegle]
            linje["afvigelse"], linje["tilladt"], linje["status"] = _vurder(
                kontrol, linje["aleta"], linje["rapport"], tolerance, _tal(a["commitment"]))
            if kontrol == "NAV":
                linje.update(valuta=f["valuta"], aleta_fv=linje["aleta"],
                             rapport_fv=linje["rapport"], afvigelse_fv=linje["afvigelse"])
            elif noegle and k:
                aleta_fv = None if linje["aleta"] is None else linje["aleta"] / k
                rapport_fv = _tal(f["n"].get(noegle))
                linje.update(valuta=f["valuta"], kurs=_tal(k), kursdato=kdato,
                             aleta_fv=aleta_fv, rapport_fv=rapport_fv,
                             afvigelse_fv=None if None in (rapport_fv, aleta_fv)
                             else rapport_fv - aleta_fv)
            ud.append(linje)
    return ud


def kontrolstatus(fonde, aleta, dato, kurser_aleta):
    """{fund_code: {kontrol: status}} for rapportens fonde, fra afstem(); til sliden med
    nøgletallene afstemt med Aleta."""
    ud = {}
    for x in afstem(aleta, {"fonde": fonde}, dato, kurser_aleta):
        ud.setdefault(x["fond"], {})[x["kontrol"]] = x["status"]
    return ud


def paa_fondene(fonde, investor, as_of, kurs):
    """Sæt f["aleta"] (irr_sammenligning()) og f["aleta_kontrol"] (kontrolstatus()) på hver fond,
    Aleta har pr. as_of. Uden et udtræk på Aletas dato ændres intet. kurs(valuta, dato) er
    data.investor_rapports kursopslag."""
    filer = find(investor, as_of)
    if filer is None:
        return
    dato = aleta_dato(as_of)
    aleta = laes_aleta(filer[0])
    raekker = {FONDE.get(a["securityCode"], a["securityCode"]): a for a in aleta}
    kurser_aleta = {}
    for v in {f["valuta"] for f in fonde}:
        k = 1.0 if v == "DKK" else kurs(v, dato)
        kurser_aleta[v] = None if k is None else (k, dato)
    alle_tx = (afstem_transaktioner(laes_transaktioner(filer[1]), {"fonde": fonde}, as_of)
               if filer[1] else [])
    for f in fonde:
        a = raekker.get(f["fund_code"])
        if a is None or a.get("irr") is None:
            continue
        tx = [t for t in alle_tx if t["fond"] == f["fund_code"]]
        f["aleta"] = irr_sammenligning(f, a, tx, dato,
                                       1.0 if f["valuta"] == "DKK" else kurs(f["valuta"], dato))
    # efter f["aleta"], så IRR-kontrollen sammenligner rapportens IRR i DKK
    status = kontrolstatus(fonde, aleta, dato, kurser_aleta)
    for f in fonde:
        if f["fund_code"] in status:
            f["aleta_kontrol"] = status[f["fund_code"]]
